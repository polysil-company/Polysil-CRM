"""The import template (FS-068): built from the fields, read back by our own
reader with every column matched, its code columns text-formatted in the saved
file, and its examples well-formed but fake. Plus rule 6 (a remembered match
never takes another field's own column) and rule 7 (a header with no rows is
refused)."""

from __future__ import annotations

import io

import openpyxl
import pytest

from api.domain.imports.fields import FIELDS, KINDS, header_score, suggest
from api.domain.imports.reader import ImportUnreadableError, SheetRow, read_upload
from api.domain.imports.rows import (
    PartnerLookups,
    Place,
    ProductLookups,
    finish_partners,
    parse_lines,
    parse_partners,
    parse_price_items,
)
from api.domain.imports.template import FILE_NAMES, HOW_TO, SHEET_TITLES, needed, template


def _book(kind: str) -> openpyxl.Workbook:
    # not read_only: a read-only sheet has no column_dimensions
    return openpyxl.load_workbook(io.BytesIO(template(kind)))


def _filled(kind: str, rows: list[list[object]], *, above: list[object] | None = None) -> bytes:
    wb = _book(kind)
    ws = wb.worksheets[0]
    if above is not None:
        ws.insert_rows(1)
        for c, value in enumerate(above, start=1):
            ws.cell(row=1, column=c, value=value)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _empty(rows: list[SheetRow], width: int) -> set[int]:
    return {c for c in range(width) if not any(r.cells[c] for r in rows)}


def _examples(kind: str) -> list[str]:
    return [f.example for f in FIELDS[kind]]


@pytest.mark.parametrize("kind", KINDS)
def test_sheet_one_is_the_labels_and_nothing_else(kind: str) -> None:
    wb = _book(kind)
    ws = wb.worksheets[0]
    assert ws.title == SHEET_TITLES[kind]
    assert wb.sheetnames == [SHEET_TITLES[kind], HOW_TO]
    assert wb.active is not None and wb.active.title == SHEET_TITLES[kind]
    assert [c.value for c in ws[1]] == [f.label for f in FIELDS[kind]]
    assert ws.max_row == 1, "an example row left in would be imported"
    assert ws.freeze_panes == "A2"
    assert FILE_NAMES[kind].endswith(".xlsx")


@pytest.mark.parametrize("kind", KINDS)
def test_code_columns_are_text_in_the_saved_file(kind: str) -> None:
    ws = _book(kind).worksheets[0]
    styled = {i for i in range(len(FIELDS[kind]))
              if ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].number_format == "@"}
    assert styled == {i for i, f in enumerate(FIELDS[kind]) if f.text}
    assert styled, "every kind has a code column"


@pytest.mark.parametrize("kind", KINDS)
def test_how_to_fill_has_every_field(kind: str) -> None:
    how = _book(kind)[HOW_TO]
    assert "Do not add rows above it" in str(how["A1"].value)
    assert [c.value for c in how[3]] == ["Column", "Needed", "What to put", "Example"]
    got = [[c.value for c in row] for row in how.iter_rows(min_row=4)]
    assert got == [[f.label, needed(kind, f.key), f.help, f.example] for f in FIELDS[kind]]
    assert all(f.help and f.example for f in FIELDS[kind]), "no blank cell on sheet 2"


def test_needed_reads_the_rules() -> None:
    assert needed("price_list_items", "rate") == "Yes"
    assert needed("price_list_items", "item_code") == "Item code or Description"
    assert needed("partners", "gstin") == "No"
    assert "GST registered" in next(f.help for f in FIELDS["partners"] if f.key == "gstin")


@pytest.mark.parametrize("kind", KINDS)
def test_a_filled_template_matches_every_column(kind: str) -> None:
    sheet = read_upload(_filled(kind, [_examples(kind)]), "t.xlsx", score=header_score(kind))
    assert sheet.header_row == 1 and len(sheet.rows) == 1
    got = suggest(kind, sheet.headers, empty=_empty(sheet.rows, len(sheet.headers)))
    assert got == {f.key: i for i, f in enumerate(FIELDS[kind])}


@pytest.mark.parametrize("kind", KINDS)
def test_only_the_needed_columns_filled(kind: str) -> None:
    # the person fills what they must; Item code stays blank, Description carries the name
    keep = {f.key for f in FIELDS[kind] if f.required} | (
        {"description"} if kind != "partners" else set())
    row = [f.example if f.key in keep else None for f in FIELDS[kind]]
    sheet = read_upload(_filled(kind, [row]), "t.xlsx", score=header_score(kind))
    got = suggest(kind, sheet.headers, empty=_empty(sheet.rows, len(sheet.headers)))
    assert got == {f.key: i for i, f in enumerate(FIELDS[kind]) if f.key in keep}


def test_partner_examples_are_well_formed_and_the_parent_resolves_to_nothing() -> None:
    # format only: "fake" is the EXAMPLE- prefix, checked below against the masters
    gj, ahd = Place("gj", "Gujarat"), Place("ahd", "Ahmedabad")
    lk = PartnerLookups(states={"gujarat": gj}, districts={("gj", "ahmedabad"): ahd})
    mapping = {f.key: i for i, f in enumerate(FIELDS["partners"])}
    rows = [SheetRow(row=2, cells=tuple(_examples("partners")))]
    results = parse_partners(rows, mapping, lk)
    assert results[0].errors == {}
    finish_partners(results, taken=set(), outside=set(), existing={})
    assert set(results[0].errors) == {"parent_code"}, "a copied example row cannot be saved"


@pytest.mark.parametrize(("kind", "parse"), [("price_list_items", parse_price_items),
                                             ("lines", parse_lines)])
def test_product_examples_are_well_formed(kind: str, parse: object) -> None:
    mapping = {f.key: i for i, f in enumerate(FIELDS[kind])}
    rows = [SheetRow(row=2, cells=tuple(_examples(kind)))]
    result = parse(rows, mapping, ProductLookups.build([]))[0]  # type: ignore[operator]
    assert set(result.errors) == {"item_code"}, result.errors


def test_example_codes_are_marked_as_examples() -> None:
    # code review F-5: a copied example must never name a real product or partner
    for kind in KINDS:
        for f in FIELDS[kind]:
            if f.key in ("item_code", "code", "parent_code", "description"):
                assert f.example.upper().startswith("EXAMPLE"), (kind, f.key)


# ── rule 6 ───────────────────────────────────────────────────────────────────

def test_a_remembered_wrong_match_does_not_take_the_templates_column() -> None:
    # 16 Oct: names under "Item Code", matched to Description and saved
    headers = ["Item code", "Description", "Rate"]
    assert suggest("price_list_items", headers, {"description": "Item Code"}) == {
        "item_code": 0, "description": 1, "rate": 2}


def test_a_remembered_alias_still_counts() -> None:
    # "item" is an alias of description, not its key or label
    assert suggest("price_list_items", ["Item", "Rate"], {"item_code": "Item"}) == {
        "item_code": 0, "rate": 1}
    assert suggest("partners", ["Distributor Code"], {"parent_code": "Distributor Code"}) == {
        "parent_code": 0}


def test_the_clients_own_layout_keeps_its_remembered_match() -> None:
    # code review F-1: no Description column, so names under "Item Code" stay names
    assert suggest("price_list_items", ["Item Code", "Rate"], {"description": "Item Code"}) == {
        "description": 0, "rate": 1}
    assert suggest("partners", ["Type", "Partner Code", "Firm", "State", "Code"],
                   {"code": "Partner Code", "parent_code": "Code"})["parent_code"] == 4
    assert suggest("partners", ["Firm", "Name"], {"contact_name": "Name"}) == {
        "contact_name": 1, "name": 0}


def test_a_contact_under_name_beside_a_contact_column_is_not_remembered() -> None:
    # the price of rule 6: with "Contact name" present, "Name" is the name field's own
    assert suggest("partners", ["Firm", "Name", "Contact name"], {"contact_name": "Name"}) == {
        "name": 1, "contact_name": 2}


# ── rule 7 ───────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("kind", KINDS)
def test_the_blank_template_is_refused(kind: str) -> None:
    with pytest.raises(ImportUnreadableError, match="no rows under them") as exc:
        read_upload(template(kind), "t.xlsx", score=header_score(kind))
    assert exc.value.code == "import_unreadable"


def test_a_chosen_header_row_with_nothing_under_it() -> None:
    with pytest.raises(ImportUnreadableError, match="Row 1 has no rows under it") as exc:
        read_upload(template("lines"), "t.xlsx", header_row=1)
    assert exc.value.code == "validation_error"


def test_a_title_row_above_the_header_renames_the_rate() -> None:
    # sheet 2 says not to; this pins what happens when someone does (EC-6)
    data = _filled("price_list_items", [["", "EXAMPLE PIPE 16 MM", 10]],
                   above=["Price list", None, "Gujarat"])
    sheet = read_upload(data, "t.xlsx", score=header_score("price_list_items"))
    assert sheet.header_row == 2
    assert sheet.headers[2] == "Rate (Gujarat)"
    assert "rate" not in suggest("price_list_items", sheet.headers)
