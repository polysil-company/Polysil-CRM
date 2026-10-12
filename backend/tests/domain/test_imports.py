"""Excel import, the pure half (FS-062): reading files, cell text, column guesses,
the mapping check and each kind's row rules. Edge cases EC-8, EC-9, EC-12, EC-15,
EC-16, EC-17 and EC-20 are asserted here."""

from __future__ import annotations

import codecs
import datetime as dt
import io
import zipfile
from decimal import Decimal

import openpyxl
import pytest

from api.domain.imports.cells import CellError, cell_text, clean_text, to_bool, to_decimal
from api.domain.imports.fields import FIELDS, mapping_problems, suggest
from api.domain.imports.reader import ImportUnreadableError, SheetRow, read_upload
from api.domain.imports.rows import (
    PartnerLookups,
    Place,
    ProductLookups,
    ProductRef,
    finish_partners,
    parse_lines,
    parse_partners,
    parse_price_items,
)


def _xlsx(rows: list[list[object]], *, formats: dict[str, str] | None = None,
          hidden_first: bool = False) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    if hidden_first:
        ws.title = "Hidden"
        ws.append(["nothing"])
        ws.sheet_state = "hidden"
        ws = wb.create_sheet("Data")
    for r in rows:
        ws.append(r)
    for ref, fmt in (formats or {}).items():
        ws[ref].number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── reading ──────────────────────────────────────────────────────────────────

def test_an_xlsx_gives_headers_and_text_rows_with_sheet_row_numbers() -> None:
    data = _xlsx([["Dealer Name", "Mob", "Rate"], [], ["Shree Agro", 9876543210, 1234.5],
                  ["Patel", "098765 43210", 10]])
    sheet = read_upload(data, "dealers.xlsx")
    assert sheet.headers == ["Dealer Name", "Mob", "Rate"]
    assert sheet.header_row == 1 and sheet.sheet == "Sheet"
    assert [r.row for r in sheet.rows] == [3, 4], "the blank row is skipped, numbers kept"
    assert sheet.rows[0].cells == ("Shree Agro", "9876543210", "1234.5")


def test_float_noise_and_percent_cells_read_as_shown() -> None:
    """EC-9: 0.1 + 0.2 is stored as 0.30000000000000004; a 10% cell holds 0.1."""
    data = _xlsx([["rate", "disc", "big"], [0.1 + 0.2, 0.1, 1e16]], formats={"B2": "0%"})
    cells = read_upload(data, "x.xlsx").rows[0].cells
    assert cells == ("0.3", "10", "10000000000000000")


def test_the_first_visible_sheet_is_the_default() -> None:
    """EC-20: a hidden first sheet is skipped."""
    sheet = read_upload(_xlsx([["code"], ["A1"]], hidden_first=True), "x.xlsx")
    assert sheet.sheet == "Data" and sheet.sheets == ["Hidden", "Data"]
    with pytest.raises(ImportUnreadableError):
        read_upload(_xlsx([["code"]]), "x.xlsx", sheet="Nope")


def test_blank_and_repeated_headers_get_names() -> None:
    sheet = read_upload(_xlsx([["Name", None, "name", "Rate"], ["a", "b", "c", "1"]]), "x.xlsx")
    assert sheet.headers == ["Name", "Column B", "name (2)", "Rate"]


def test_an_explicit_header_row_skips_title_rows() -> None:
    data = _xlsx([["Polysil price list 2026"], [], ["code", "rate"], ["P1", 10]])
    sheet = read_upload(data, "x.xlsx", header_row=3)
    assert sheet.headers == ["code", "rate"] and sheet.rows[0].row == 4
    with pytest.raises(ImportUnreadableError):
        read_upload(data, "x.xlsx", header_row=2)


def test_an_unformulated_formula_reads_as_empty() -> None:
    """EC-16: openpyxl writes no cached result, so the cell is empty; the service turns
    a wholly empty mapped column into one message."""
    sheet = read_upload(_xlsx([["a", "b"], [2, "=A2*2"]]), "x.xlsx")
    assert sheet.rows[0].cells == ("2", "")


def test_an_inflated_zip_is_refused_before_it_is_opened() -> None:
    """EC-8: the directory's sizes are summed first."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/sharedStrings.xml", b"\0" * (101 * 1024 * 1024))
    data = buf.getvalue()
    assert len(data) < 1024 * 1024
    with pytest.raises(ImportUnreadableError, match="100 MB"):
        read_upload(data, "bomb.xlsx")


@pytest.mark.parametrize(("data", "name"), [
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "old.xls"),
    (b"PK\x03\x04garbage", "x.xlsx"),
    (b"not really", "x.xlsx"),
    (b"\x00\x01\x02binary", "x.csv")])
def test_unreadable_files_are_refused(data: bytes, name: str) -> None:
    with pytest.raises(ImportUnreadableError):
        read_upload(data, name)


def test_a_zip_that_is_not_a_workbook_is_refused() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("readme.txt", "hi")
    with pytest.raises(ImportUnreadableError, match="not an Excel"):
        read_upload(buf.getvalue(), "x.xlsx")


@pytest.mark.parametrize("encoded", [
    codecs.BOM_UTF8 + b"code,name\nA1,Shree Agro\n",
    codecs.BOM_UTF16_LE + "code\tname\r\nA1\tShree Agro\r\n".encode("utf-16-le"),
    "code;name\nA1;Shree Agro\n".encode("cp1252"),
])
def test_csv_in_common_encodings_and_separators(encoded: bytes) -> None:
    """EC-15."""
    sheet = read_upload(encoded, "x.csv")
    assert sheet.headers == ["code", "name"]
    assert sheet.rows[0].cells == ("A1", "Shree Agro")


def test_a_gujarati_csv_in_utf8_round_trips() -> None:
    sheet = read_upload("નામ,code\nરાજકોટ,A1\n".encode(), "x.csv")
    assert sheet.headers[0] == "નામ" and sheet.rows[0].cells[0] == "રાજકોટ"


def test_too_many_rows_is_its_own_code() -> None:
    body = "code\n" + "".join(f"P{i}\n" for i in range(5001))
    with pytest.raises(ImportUnreadableError) as exc:
        read_upload(body.encode(), "x.csv")
    assert exc.value.code == "import_too_many_rows"
    assert len(read_upload(("code\n" + "".join(f"P{i}\n" for i in range(5000))).encode(),
                           "x.csv").rows) == 5000


# ── cells ────────────────────────────────────────────────────────────────────

def test_cell_text_forms() -> None:
    assert cell_text(None) == ""
    assert cell_text(9876543210.0) == "9876543210"
    assert cell_text(1234.4999999999998) == "1234.5"
    assert cell_text(dt.datetime(2026, 10, 12)) == "2026-10-12"
    assert cell_text(True) == "TRUE"
    nbsp, zwsp = chr(0xA0), chr(0x200B)
    assert clean_text(f"{nbsp}Shree{zwsp} Agro{nbsp}") == "Shree Agro"


@pytest.mark.parametrize(("text", "want"), [
    ("1,50,000", "150000.00"), ("₹1,50,000.5", "150000.50"), ("Rs. 99", "99.00"),
    ("12.30", "12.30")])
def test_money_accepts_indian_forms(text: str, want: str) -> None:
    assert str(to_decimal(text, places=2)) == want


@pytest.mark.parametrize("text", ["12.345", "abc", "NaN", "1e400x"])
def test_money_refuses_more_decimals_and_non_numbers(text: str) -> None:
    with pytest.raises(CellError):
        to_decimal(text, places=2)


def test_yes_and_no() -> None:
    assert to_bool("Y") is True and to_bool("no") is False
    with pytest.raises(CellError):
        to_bool("maybe")


# ── fields and mapping ───────────────────────────────────────────────────────

def test_suggestions_from_built_in_names() -> None:
    headers = ["Dealer Name", "Mob", "GST No", "Dist", "State", "Type", "Dealer Code"]
    got = suggest("partners", headers)
    assert got == {"name": 0, "mobile": 1, "gstin": 2, "district": 3, "state": 4,
                   "partner_type": 5, "code": 6}


def test_a_remembered_mapping_wins_and_columns_are_used_once() -> None:
    headers = ["Firm", "Contact", "Phone"]
    got = suggest("partners", headers, {"name": "Firm", "mobile": "Contact"})
    assert got["name"] == 0 and got["mobile"] == 1
    assert "contact_name" not in got, "column 1 is already the mobile"


def test_mapping_problems() -> None:
    assert mapping_problems("lines", {"qty": 0, "item_code": 1}, 2) == {}
    problems = mapping_problems("lines", {"qty": 0, "description": 0, "bogus": 1}, 2)
    assert problems["mapping.description"].startswith("that column is already")
    assert problems["mapping.bogus"] == "not a field of this import"
    assert mapping_problems("lines", {"qty": 0}, 2) == {
        "mapping.item_code": "match item_code or description"}
    assert mapping_problems("lines", {"qty": 5, "item_code": 0}, 2) == {
        "mapping.qty": "no column 5"}
    assert "mapping.rate" in mapping_problems("price_list_items", {"item_code": 0}, 2)


def test_every_field_has_a_label() -> None:
    for kind, specs in FIELDS.items():
        assert len({s.key for s in specs}) == len(specs), kind
        assert all(s.label for s in specs), kind


# ── partners ─────────────────────────────────────────────────────────────────

GJ = Place("gj", "Gujarat")
RAJKOT = Place("rjt", "Rajkot")
LK = PartnerLookups(states={"gujarat": GJ, "gj": GJ}, districts={("gj", "rajkot"): RAJKOT})
COLS = ["partner_type", "code", "name", "state", "district", "parent_code", "mobile", "gstin",
        "credit_limit", "is_gst_registered"]
MAP = {k: i for i, k in enumerate(COLS)}


def _rows(*cells: tuple[str, ...]) -> list[SheetRow]:
    return [SheetRow(row=i + 2, cells=c) for i, c in enumerate(cells)]


def test_a_good_partner_row_is_normalised() -> None:
    [res] = parse_partners(_rows(("Dealer", "D-1", "Shree Agro", "GJ", "rajkot.", "",
                                  "098765 43210", "24aaaaa0000a1z5", "1,50,000", "")), MAP, LK)
    assert res.ok, res.errors
    assert res.values["partner_type"] == "dealer" and res.values["territory_id"] == "rjt"
    assert res.values["territory"] == "Rajkot, Gujarat"
    assert res.values["mobile"] == "919876543210"
    assert res.values["gstin"] == "24AAAAA0000A1Z5" and res.values["is_gst_registered"] is True
    assert res.values["credit_limit"] == "150000.00"


def test_every_error_in_a_row_is_reported_at_once() -> None:
    [res] = parse_partners(_rows(("wholesaler", "", "", "Kerala", "", "", "12345", "BAD", "-5",
                                  "yes")), MAP, LK)
    assert set(res.errors) == {"partner_type", "code", "name", "state", "mobile", "gstin",
                               "credit_limit"}


def test_registered_without_a_gstin_is_an_error() -> None:
    [res] = parse_partners(_rows(("dealer", "D1", "A", "GJ", "", "", "", "", "", "Y")), MAP, LK)
    assert res.errors == {"gstin": "required when GST registered"}


def test_a_district_must_be_in_its_state() -> None:
    [res] = parse_partners(_rows(("dealer", "D1", "A", "GJ", "Rajkott", "", "", "", "", "")),
                           MAP, LK)
    assert res.errors == {"district": "no district 'Rajkott' in Gujarat"}


def test_duplicate_codes_in_a_file_ignore_case_and_mark_both_rows() -> None:
    out = parse_partners(_rows(("dealer", "d-101", "A", "GJ", "", "", "", "", "", ""),
                               ("dealer", "D-101", "B", "GJ", "", "", "", "", "", "")), MAP, LK)
    assert out[0].errors == {"code": "also on row 3"}
    assert out[1].errors == {"code": "also on row 2"}


def test_taken_codes_and_territories_outside_the_area() -> None:
    """EC-10 and EC-11."""
    out = parse_partners(_rows(("dealer", "D1", "A", "GJ", "Rajkot", "", "", "", "", "")), MAP, LK)
    finish_partners(out, taken={"d1"}, outside={"rjt"}, existing={})
    assert out[0].errors == {"code": "D1 is already in use",
                             "district": "Rajkot, Gujarat is outside your area"}


def test_parents_in_the_file_and_the_cascade() -> None:
    """EC-12: a dealer under a distributor with errors gets an error of its own, and so
    does its sub-dealer, whatever the order of the rows."""
    out = parse_partners(_rows(
        ("sub dealer", "S1", "Sub", "GJ", "", "D1", "", "", "", ""),
        ("dealer", "D1", "Dealer", "GJ", "", "X1", "", "", "", ""),
        ("distributor", "X1", "Dist", "GJ", "", "", "", "BADGST", "", ""),
        ("dealer", "D2", "Dealer 2", "GJ", "", "X9", "", "", "", ""),
        ("dealer", "D3", "Dealer 3", "GJ", "", "OLD", "", "", "", ""),
        ("dealer", "D4", "Dealer 4", "GJ", "", "D3", "", "", "", ""),
    ), MAP, LK)
    finish_partners(out, taken=set(), outside=set(),
                    existing={"old": ("old-id", "distributor")})
    by_code = {r.values["code"]: r for r in out}
    assert by_code["X1"].errors == {"gstin": "not a GSTIN"}
    assert by_code["D1"].errors == {"parent_code": "the parent on row 4 has errors"}
    assert by_code["S1"].errors == {"parent_code": "the parent on row 3 has errors"}
    assert by_code["D2"].errors == {"parent_code": "no partner with code X9"}
    assert by_code["D3"].ok and by_code["D3"].values["parent_id"] == "old-id"
    assert by_code["D4"].errors["parent_code"].startswith("wrong type")


def test_a_skipped_parent_strands_its_children() -> None:
    out = parse_partners(_rows(("distributor", "X1", "Dist", "GJ", "", "", "", "", "", ""),
                               ("dealer", "D1", "Dealer", "GJ", "", "x1", "", "", "", "")), MAP, LK)
    finish_partners(out, taken=set(), outside=set(), existing={})
    assert out[1].ok and out[1].values["parent_row"] == 2
    again = parse_partners(_rows(("distributor", "X1", "Dist", "GJ", "", "", "", "", "", ""),
                                 ("dealer", "D1", "Dealer", "GJ", "", "x1", "", "", "", "")),
                           MAP, LK)
    finish_partners(again, taken=set(), outside=set(), existing={}, skipped=frozenset({2}))
    assert again[1].errors == {"parent_code": "the parent on row 2 has errors"}


def test_a_partner_cannot_be_its_own_parent() -> None:
    out = parse_partners(_rows(("dealer", "D1", "A", "GJ", "", "d1", "", "", "", "")), MAP, LK)
    finish_partners(out, taken=set(), outside=set(), existing={})
    assert out[0].errors == {"parent_code": "a partner cannot be its own parent"}


# ── products ─────────────────────────────────────────────────────────────────

PIPE = ProductRef("p1", "PIPE 16MM", True, "MTR", 0, Decimal(100))
VALVE = ProductRef("p2", "VALVE 2IN", True, "NOS", 0, None)
OLD = ProductRef("p3", "OLD FILTER", False, "NOS", 0, None)
WIRE = ProductRef("p4", "WIRE", True, "KG", 2, None)
PL = ProductLookups(by_code={"c1": PIPE, "c2": VALVE, "c3": OLD, "c4": WIRE},
                    by_description={"pipe 16mm": PIPE, "valve 2in": VALVE, "old filter": OLD,
                                    "wire": WIRE})


def test_price_rows_match_by_code_or_description() -> None:
    m = {"item_code": 0, "description": 1, "rate": 2}
    out = parse_price_items(_rows(("C1", "", "10.50"), ("", "valve 2in", "99"),
                                  ("C1", "VALVE 2IN", "5"), ("C9", "", "5"),
                                  ("C4", "some other words", "12.345"), ("", "", "1")), m, PL)
    assert out[0].values == {"product_id": "p1", "product": "PIPE 16MM", "rate": "10.50"}
    assert out[1].ok and out[1].values["product_id"] == "p2"
    assert out[2].errors["description"] == "names a different product from the item code"
    assert out[3].errors == {"item_code": "no product with item code C9"}
    assert out[4].errors == {"rate": "at most 2 decimal places"}, "unknown wording is ignored"
    assert out[5].errors == {"item_code": "required"}


def test_a_product_twice_in_a_price_file_is_an_error_on_both() -> None:
    out = parse_price_items(_rows(("C1", "1"), ("c1", "2")), {"item_code": 0, "rate": 1}, PL)
    assert out[0].errors == {"item_code": "also on row 3"}
    assert out[1].errors == {"item_code": "also on row 2"}


def test_line_rows_apply_the_units_rules() -> None:
    """EC-17: the save refuses these, so the check does too."""
    m = {"item_code": 0, "qty": 1, "discount_pct": 2}
    out = parse_lines(_rows(("C1", "250", ""), ("C1", "300", "10"), ("C2", "1.5", ""),
                            ("C3", "1", ""), ("C4", "2.25", "120"), ("C2", "0", "")), m, PL)
    assert out[0].errors == {"qty": "sold in multiples of 100 MTR"}
    assert out[1].ok and out[1].values["qty"] == "300" and out[1].values["discount_pct"] == "10.000"
    assert out[2].errors == {"qty": "NOS takes 0 decimal places"}
    assert out[3].errors == {"item_code": "this product is no longer sold"}
    assert out[4].errors == {"discount_pct": "must be at most 100"}
    assert out[4].values["qty"] == "2.25"
    assert out[5].errors == {"qty": "must be more than 0"}
