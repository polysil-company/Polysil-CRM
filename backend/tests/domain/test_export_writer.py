"""The FS-030 workbook writer, without a database: typed cells, IST, and the
three ways a text value can break a workbook (review findings 2, 3, 7, 8)."""



from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel

from api.services import exports
from api.services.exports import Column

COLS = (Column("Name", "name"), Column("Amount", "amount", "money"),
        Column("At", "at", "datetime"), Column("On", "on", "date"),
        Column("Owner", "owner.full_name"), Column("Crops", "crops", "names"),
        Column("Active", "active", "bool"), Column("Count", "n", "int"),
        Column("Mobile", "mobile"))


def _sheet(rows: list[dict[str, Any]]) -> Any:
    wb = load_workbook(io.BytesIO(exports.workbook("Test", COLS, rows)))
    return wb.active


def test_header_is_frozen_and_filtered_and_an_empty_list_is_header_only() -> None:
    ws = _sheet([])
    assert [c.value for c in ws[1]] == [c.header for c in COLS]
    assert ws.max_row == 1
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")


def test_money_is_a_number_cell_never_text() -> None:
    ws = _sheet([{"amount": "123456.78"}])
    cell = ws["B2"]
    assert cell.data_type == "n"
    assert Decimal(str(cell.value)) == Decimal("123456.78")
    assert cell.number_format == "#,##0.00"


@pytest.mark.parametrize("value", ["=1+1", "+919876543210", "-2+3", "@SUM(A1)",
                                   '=HYPERLINK("http://x","y")'])
def test_text_that_looks_like_a_formula_stays_text(value: str) -> None:
    ws = _sheet([{"name": value, "mobile": value}])
    for ref in ("A2", "I2"):
        assert ws[ref].data_type == "s", ref
        assert ws[ref].value == value


def test_a_utc_timestamp_lands_as_the_ist_wall_time() -> None:
    ws = _sheet([{"at": "2026-10-03T20:00:00+00:00", "on": "2026-10-04"}])
    assert ws["C2"].value == dt.datetime(2026, 10, 4, 1, 30)
    assert ws["D2"].value == dt.datetime(2026, 10, 4, 0, 0)


def test_control_characters_are_stripped_and_long_text_is_cut() -> None:
    long = "x" * (exports.CELL_MAX + 50)
    ws = _sheet([{"name": "bad\x01char\x1f", "mobile": long}])
    assert ws["A2"].value == "badchar"
    assert len(ws["I2"].value) == exports.CELL_MAX
    assert ws["I2"].value.endswith("…")


def test_nulls_are_empty_cells_and_nested_lists_and_bools_flatten() -> None:
    ws = _sheet([{"owner": None, "crops": [{"name": "Cotton"}, {"name": "Castor"}],
                  "active": False, "n": 7, "amount": None}])
    assert ws["E2"].value is None
    assert ws["B2"].value is None
    assert ws["F2"].value == "Cotton, Castor"
    assert ws["G2"].value == "No"
    assert ws["H2"].value == 7
    ws = _sheet([{"owner": {"id": "u", "full_name": "Ravi Desai"}}])
    assert ws["E2"].value == "Ravi Desai"


def test_numbers_held_as_text_keep_their_leading_zeros() -> None:
    ws = _sheet([{"mobile": "0012345"}])
    assert ws["I2"].value == "0012345" and ws["I2"].data_type == "s"


class _Item(BaseModel):
    n: int


async def test_drain_walks_every_page_and_refuses_one_past_the_cap() -> None:
    pages = {None: ([_Item(n=1), _Item(n=2)], "c1"), "c1": ([_Item(n=3)], None)}
    calls: list[str | None] = []

    async def fetch(cursor: str | None) -> tuple[list[_Item], str | None]:
        calls.append(cursor)
        return pages[cursor]

    rows = await exports.drain(fetch, cap=3)
    assert [r["n"] for r in rows] == [1, 2, 3] and calls == [None, "c1"]
    with pytest.raises(exports.ExportTooLargeError):
        await exports.drain(fetch, cap=2)


def test_filename_is_the_ist_date() -> None:
    assert exports.filename("leads", dt.date(2026, 10, 3)) == "leads-2026-10-03.xlsx"


def test_every_column_path_names_a_field_of_its_list_item() -> None:
    """A renamed schema field must fail here, not leave an empty column."""
    from api.schemas import (
        complaints,
        leads,
        masters,
        orders,
        quotations,
        subsidy_applications,
        tasks,
        users,
    )

    def fields_of(model: type[BaseModel], path: str) -> bool:
        current: Any = model
        for part in path.split("."):
            if not (isinstance(current, type) and issubclass(current, BaseModel)):
                return False
            info = current.model_fields.get(part)
            if info is None:
                return False
            ann = info.annotation
            args = [a for a in getattr(ann, "__args__", ()) if a is not type(None)]
            current = args[0] if args else ann
        return True

    pairs = [(leads.Lead, exports.LEADS), (quotations.QuotationSummary, exports.QUOTATIONS),
             (orders.OrderSummary, exports.ORDERS),
             (complaints.ComplaintSummary, exports.COMPLAINTS), (tasks.Task, exports.TASKS),
             (subsidy_applications.Application, exports.SUBSIDY_APPLICATIONS),
             (masters.Partner, exports.PARTNERS), (users.UserRow, exports.USERS)]
    for model, cols in pairs:
        for col in cols:
            assert fields_of(model, col.path), f"{model.__name__}: {col.path}"


def test_the_pims_sheet_keeps_a_formula_looking_name_as_text() -> None:
    """ISS-200: the PIMS sheet predates the writer and appended raw strings."""
    from api.services.subsidy_applications import PIMS_COLUMNS, pims_workbook

    row = ["=HYPERLINK(\"http://x\",\"y\")"] + [None] * (len(PIMS_COLUMNS) - 1)
    ws = load_workbook(io.BytesIO(pims_workbook([tuple(row)]))).active
    assert ws["A2"].data_type == "s" and ws["A2"].value == row[0]


def test_an_export_writes_one_log_line_with_filter_names_never_values() -> None:
    """Code review F-1: the line must reach structlog, which production configures."""
    from structlog.testing import capture_logs

    with capture_logs() as logs:
        exports.log_export("u1", "leads", {"q": "9876543210", "stage": "new", "owner": None}, 3)
    assert len(logs) == 1
    entry = logs[0]
    assert entry["event"] == "export" and entry["list"] == "leads" and entry["rows"] == 3
    assert entry["filters"] == ["q", "stage"]
    assert "9876543210" not in repr(entry)
