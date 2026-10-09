"""/subsidy-applications, end to end over ASGI (FS-009).

The order tests' shop, plus a State Co-ordinator on the shop's state and an item
code on its product. The engine is replaced by a fixed calculation of two crops,
so every rule here runs without the client's masters; one test at the end runs
the real engine against the Drip sample and is skipped where the samples are
absent, as FS-008's tests are.
"""

# ruff: noqa: E501  (embedded SQL, kept on one line so each statement reads whole)

from __future__ import annotations

import datetime as dt
import io
import uuid
from collections.abc import AsyncIterator, Callable
from dataclasses import dataclass
from decimal import Decimal
from typing import Any

import httpx
import pytest
import pytest_asyncio
from argon2 import PasswordHasher
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from api.authz.predicate import Caller
from api.schemas import subsidy_applications as sch
from api.schemas.subsidy import CalculateRequest, CalculateResponse
from api.services import subsidy_applications as service
from api.services.clock import today_ist
from api.storage import UnconfiguredStorage
from tests.api import test_complaints as complaints
from tests.api import test_order_concurrency as conc
from tests.api import test_order_endpoints as endpoints
from tests.api.conftest import PASSWORD, V1, _key, _login
from tests.domain.subsidy_masters import needs_samples

pytestmark = pytest.mark.db

shop = endpoints.shop
Shop = endpoints.Shop
Sessions = Callable[[], AsyncSession]
APPS = f"{V1}/subsidy-applications"

JPEG = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00" + b"\x01" * 64
PDF = b"%PDF-1.7\n" + b"\x02" * 64


@dataclass
class Office:
    shop: Shop
    coordinator: str        # email of a State Co-ordinator on the shop's state
    coordinator_id: str
    stranger: str           # email of a State Co-ordinator with no territory
    stranger_id: str
    item_code: str


@pytest_asyncio.fixture
async def office(shop: Shop, sessions: Sessions) -> AsyncIterator[Office]:
    tag = uuid.uuid4().hex[:8]
    s = sessions()
    hasher = PasswordHasher()
    ids: dict[str, str] = {}
    for who in ("sc", "stranger"):
        ids[who] = str((await s.execute(text(
            "INSERT INTO app_user (user_type, email, password_hash, full_name, role_id, org_unit_id) "
            "SELECT 'staff', :e, :p, :n, r.id, CAST(:o AS uuid) FROM role r WHERE r.code = 'state_coordinator' "
            "RETURNING id"),
            {"e": f"sa_{who}_{tag}@polysil.in", "p": hasher.hash(PASSWORD), "n": f"Coordinator {who}",
             "o": shop.office})).scalar_one())
    await s.execute(text("INSERT INTO user_territory (user_id, territory_id) VALUES (CAST(:u AS uuid), CAST(:t AS uuid))"),
                    {"u": ids["sc"], "t": shop.state})
    item_code = f"IT{tag}".upper()
    await s.execute(text("UPDATE product SET item_code = :c WHERE id = CAST(:p AS uuid)"),
                    {"c": item_code, "p": shop.product})
    # FS-039: an application takes the scheme of the lead's state. GGRC moves to the
    # shop's state for the test and back after, whether it was on Gujarat or nowhere
    ggrc_state, ggrc_by = (await s.execute(text(
        "SELECT state_territory_id, updated_by FROM subsidy_scheme WHERE code = 'GGRC'"))).one()
    await s.execute(text("UPDATE subsidy_scheme SET state_territory_id = CAST(:s AS uuid) WHERE code = 'GGRC'"),
                    {"s": shop.state})
    await s.commit()
    await s.close()
    try:
        yield Office(shop, f"sa_sc_{tag}@polysil.in", ids["sc"], f"sa_stranger_{tag}@polysil.in",
                     ids["stranger"], item_code)
    finally:
        c = sessions()
        # updated_by too: a test that links GGRC through the API stamps its own admin
        await c.execute(text("UPDATE subsidy_scheme SET state_territory_id = :s, updated_by = :u WHERE code = 'GGRC'"),
                        {"s": ggrc_state, "u": ggrc_by})
        apps = "(SELECT id FROM subsidy_application WHERE territory_id = CAST(:d AS uuid))"
        people = [ids["sc"], ids["stranger"]]
        for stmt in (
            f"DELETE FROM subsidy_stage_value WHERE entry_id IN (SELECT id FROM subsidy_stage_entry WHERE application_id IN {apps})",
            f"DELETE FROM subsidy_stage_entry WHERE application_id IN {apps}",
            f"DELETE FROM subsidy_document WHERE application_id IN {apps}",
            f"DELETE FROM activity_event WHERE entity_type = 'subsidy_application' AND entity_id IN {apps}",
            "DELETE FROM subsidy_application WHERE territory_id = CAST(:d AS uuid)",
            "DELETE FROM subsidy_app_counter WHERE state_code = :c",
            "DELETE FROM idempotency_record WHERE user_id = ANY(CAST(:people AS uuid[]))",
            "DELETE FROM activity_event WHERE actor_id = ANY(CAST(:people AS uuid[]))",
            "DELETE FROM session WHERE user_id = ANY(CAST(:people AS uuid[]))",
            "DELETE FROM login_attempt WHERE identifier = ANY(CAST(:emails AS citext[]))",
            "DELETE FROM user_territory WHERE user_id = ANY(CAST(:people AS uuid[]))",
            "DELETE FROM app_user WHERE id = ANY(CAST(:people AS uuid[]))",
            "UPDATE product SET item_code = NULL WHERE id = CAST(:p AS uuid)",
        ):
            await c.execute(text(stmt), {"d": shop.district, "c": shop.code, "people": people,
                                         "p": shop.product,
                                         "emails": [f"sa_sc_{tag}@polysil.in", f"sa_stranger_{tag}@polysil.in"]})
        await c.commit()
        await c.close()


# ── a fixed calculation ──────────────────────────────────────────────────────

_BLOCK_KEYS = ("head_unit", "field_unit", "a_plus_b", "cgst_ab", "sgst_ab", "installation", "cgst_c",
               "sgst_c", "total_abc_gst", "insurance", "cgst_d", "sgst_d", "inspection", "cgst_e",
               "sgst_e", "education", "sump", "cost_excl_gst", "total_cgst", "total_sgst", "total_gst",
               "total_incl_gst")


def _blocks(total: str, education: str = "0") -> dict[str, str]:
    out = dict.fromkeys(_BLOCK_KEYS, "0")
    out.update(total_incl_gst=total, education=education)
    return out


def _category(code: str, subsidy: str, farmer: str, applicable: bool = True,
              reason: str | None = None) -> dict[str, Any]:
    return {"code": code, "name": code.replace("_", " ").title(), "pct": "80", "variant": "regular",
            "applicable": applicable, "reason": reason, "subsidy": subsidy, "farmer_share": farmer,
            "subsidy_pct": "80.00", "gsdma_farmer_share": None}


def _crop(area: str, categories: list[dict[str, Any]]) -> dict[str, Any]:
    return {"crop": "Cotton", "inter_crop": None, "area": area, "lateral_spacing_designed": "1.2",
            "lateral_spacing_standard": "1.2", "lateral_spacing_for_subsidy": "1.2",
            "blocks": _blocks("50000"),
            "jantri": {"regular": "1", "regular_with_sump": "1", "regular_for_cap": "1", "seven_year": None},
            "categories": categories, "warnings": []}


def _fake_response(req: CalculateRequest) -> CalculateResponse:
    """Two crops. `small_farmer` applies on both; `general` only on the first."""
    return CalculateResponse.model_validate({
        "system_type": req.system_type, "scheme": "GGRC",
        "masters": {"as_of": "2020-06-15", "formula_version": "test-1",
                    "regular_matrix_id": str(uuid.uuid4()), "seven_year_matrix_id": str(uuid.uuid4()),
                    "quantity_matrix_id": None},
        "crops": [
            _crop("1.5", [_category("small_farmer", "40000.00", "10000.00"),
                          _category("general", "30000.00", "20000.00")]),
            _crop("0.75", [_category("small_farmer", "30000.50", "7500.25"),
                           _category("general", "0", "0", False, "over the cap for this crop")]),
        ],
        "total": {"blocks": _blocks("95000.75", education="500.00")},
        "sprinkler": None, "warnings": []})


@pytest.fixture
def fake_engine(monkeypatch: pytest.MonkeyPatch) -> list[CalculateRequest]:
    seen: list[CalculateRequest] = []

    async def calculate(db: AsyncSession, req: CalculateRequest) -> CalculateResponse:
        seen.append(req)
        return _fake_response(req)
    monkeypatch.setattr(service.engine, "calculate", calculate)
    return seen


def _calc(o: Office, system: str = "drip") -> dict[str, Any]:
    return {"system_type": system, "installation_rate_per_ha": "2000",
            "crops": [{"crop": "Cotton", "area": "1.5", "lateral_spacing": "1.2",
                       "lines": [{"description": "16 mm lateral", "uom": "m", "rate": "12.50",
                                  "qty": "100", "product_id": o.shop.product.upper()}]},
                      {"crop": "Castor", "area": "0.75", "lateral_spacing": "1.2", "lines": []}],
            "head_lines": [{"description": "Screen filter", "uom": "No", "rate": "3500", "qty": "1"}]}


# ── helpers ──────────────────────────────────────────────────────────────────

async def _as(client: httpx.AsyncClient, o: Office, role: str) -> dict[str, str]:
    if role == "state_coordinator":
        return await _login(client, o.coordinator, PASSWORD)
    if role == "stranger":
        return await _login(client, o.stranger, PASSWORD)
    return await _login(client, o.shop.users[role], PASSWORD)


def _code(r: httpx.Response) -> str:
    return str(r.json()["error"]["code"])


def _fields(r: httpx.Response) -> dict[str, str]:
    return dict(r.json()["error"].get("fields") or {})


async def _lead(client: httpx.AsyncClient, o: Office, h: dict[str, str], *, inquiry: str = "subsidised",
                mis: str = "drip", to: tuple[str, ...] = ("contacted", "qualified")) -> dict[str, Any]:
    r = await client.post(f"{V1}/leads", headers={**h, **_key()}, json={
        "farmer_name": "Rameshbhai Patel", "mobile": "98" + f"{uuid.uuid4().int % 10**8:08d}",
        "territory_id": o.shop.district, "inquiry_type": inquiry, "mis_system": mis, "village": "Anand"})
    assert r.status_code == 201, r.text
    lead: dict[str, Any] = r.json()["data"]
    for stage in to:
        r = await client.post(f"{V1}/leads/{lead['id']}/transition", json={"to_stage": stage},
                              headers={**h, **_key()})
        assert r.status_code == 200, r.text
    return lead


async def _forward(client: httpx.AsyncClient, h: dict[str, str], lead_id: str, o: Office,
                   category: str = "small_farmer", system: str = "drip") -> httpx.Response:
    return await client.post(APPS, headers={**h, **_key()}, json={
        "lead_id": lead_id, "category_code": category, "calculation": _calc(o, system),
        "survey_no": "112/2"})


async def _app(client: httpx.AsyncClient, o: Office, h: dict[str, str]) -> dict[str, Any]:
    lead = await _lead(client, o, h)
    r = await _forward(client, h, lead["id"], o)
    assert r.status_code == 201, r.text
    app: dict[str, Any] = r.json()["data"]
    return app


async def _record(client: httpx.AsyncClient, h: dict[str, str], app_id: str, stage: str,
                  days_ago: int = 0, values: dict[str, Any] | None = None,
                  remark: str | None = None) -> httpx.Response:
    body: dict[str, Any] = {"stage_code": stage,
                            "occurred_on": (today_ist() - dt.timedelta(days=days_ago)).isoformat(),
                            "values": values or {}}
    if remark is not None:
        body["remark"] = remark
    return await client.post(f"{APPS}/{app_id}/stages", headers={**h, **_key()}, json=body)


def _ago(days: int) -> str:
    return (today_ist() - dt.timedelta(days=days)).isoformat()


def _days(iso: str) -> int:
    """Days from `iso` to today, read after the response: a run that crosses IST
    midnight during a request still agrees with the server."""
    return (today_ist() - dt.date.fromisoformat(iso)).days


# ── forwarding a lead (rules 1 to 4) ─────────────────────────────────────────

async def test_forwarding_a_lead_wins_it_and_stores_the_category_figures(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    lead = await _lead(client, office, fo)
    r = await _forward(client, fo, lead["id"], office)
    assert r.status_code == 201, r.text
    app = r.json()["data"]
    assert app["application_no"].startswith(f"SA/{office.shop.code}/"), app["application_no"]
    assert app["status"] == "open" and app["current_stage"]["seq"] == 4
    assert _days(app["current_stage"]["since"]) == 0
    # the per-crop figures of the chosen category, summed (review B-6)
    assert app["figures"] == {"total_cost": "95000.75", "subsidy": "70000.50", "farmer_share": "17500.25"}
    assert app["category"]["code"] == "small_farmer" and app["total_area"] == "2.250"
    assert app["farmer_name"] == "Rameshbhai Patel" and app["survey_no"] == "112/2"
    assert app["lead"]["id"] == lead["id"] and app["owner"]["id"] == office.shop.ids["field_officer"]
    since = app["current_stage"]["since"]
    assert app["reg_no"] is None and app["ageing"] == {"days_in_stage": _days(since), "days_since_inward": _days(since)}
    assert app["documents"]["uploaded"] == 0 and app["documents"]["listed"] >= 20

    r = await client.get(f"{V1}/leads/{lead['id']}", headers=fo)
    assert r.json()["data"]["stage"] == "won"
    timeline = (await client.get(f"{V1}/leads/{lead['id']}/timeline", headers=fo)).json()["data"]
    won = next(e for e in timeline if e["kind"] == "lead.stage_changed" and e["payload"].get("to") == "won")
    assert won["payload"]["via"] == "subsidy_application"
    assert any(e["kind"] == "subsidy.created" for e in timeline)

    stored = (await client.get(f"{APPS}/{app['id']}/calculation", headers=fo)).json()["data"]
    assert stored["masters"]["formula_version"] == "test-1" and len(stored["crops"]) == 2
    page = (await client.get(APPS, headers=fo, params={"q": app["application_no"]})).json()
    assert [a["id"] for a in page["data"]] == [app["id"]]
    entries = (await client.get(f"{APPS}/{app['id']}/stages", headers=fo)).json()["data"]
    assert [e["stage"]["code"] for e in entries] == ["application_in_process"]


async def test_a_lead_forwards_once_until_its_application_is_cancelled(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    lead = await _lead(client, office, fo)
    first = (await _forward(client, fo, lead["id"], office)).json()["data"]
    r = await _forward(client, fo, lead["id"], office)
    assert r.status_code == 422 and _code(r) == "already_forwarded", r.text
    r = await client.post(f"{APPS}/{first['id']}/cancel", headers={**fo, **_key()},
                          json={"reason": "  Wrong category  "})
    assert r.status_code == 200 and r.json()["data"]["status"] == "cancelled", r.text
    assert r.json()["data"]["cancellation"] == {"reason": "Wrong category"}
    # the lead is won now; a won lead still forwards (rule 1)
    r = await _forward(client, fo, lead["id"], office)
    assert r.status_code == 201, r.text
    assert r.json()["data"]["application_no"] != first["application_no"]
    r = await client.post(f"{APPS}/{first['id']}/cancel", headers={**fo, **_key()}, json={"reason": "again"})
    assert r.status_code == 409 and _code(r) == "status_changed", r.text


async def test_forwarding_refuses_a_lead_that_cannot_carry_an_application(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    commercial = await _lead(client, office, fo, inquiry="commercial")
    r = await _forward(client, fo, commercial["id"], office)
    assert r.status_code == 422 and _code(r) == "lead_not_subsidised", r.text
    fresh = await _lead(client, office, fo, to=())
    r = await _forward(client, fo, fresh["id"], office)
    assert r.status_code == 422 and _code(r) == "lead_not_forwardable", r.text
    automation = await _lead(client, office, fo, mis="automation")
    r = await _forward(client, fo, automation["id"], office)
    assert r.status_code == 422 and "calculation.system_type" in _fields(r), r.text
    lead = await _lead(client, office, fo)
    r = await _forward(client, fo, lead["id"], office, system="sprinkler")
    assert r.status_code == 422 and "calculation.system_type" in _fields(r), r.text
    r = await _forward(client, fo, lead["id"], office, category="general")
    assert r.status_code == 422 and _fields(r)["category_code"] == "over the cap for this crop", r.text
    r = await _forward(client, fo, lead["id"], office, category="no_such_category")
    assert r.status_code == 422 and "category_code" in _fields(r), r.text
    r = await client.post(APPS, headers={**fo, **_key()}, json={
        "lead_id": str(uuid.uuid4()), "category_code": "small_farmer", "calculation": _calc(office)})
    assert r.status_code == 404, r.text
    # none of the refusals moved the lead
    for refused, stage in ((commercial, "qualified"), (fresh, "new"), (lead, "qualified")):
        got = (await client.get(f"{V1}/leads/{refused['id']}", headers=fo)).json()["data"]["stage"]
        assert got == stage, (refused["id"], got)


async def test_engine_refusals_come_back_under_calculation(
        client: httpx.AsyncClient, office: Office, monkeypatch: pytest.MonkeyPatch) -> None:
    from api.errors import ValidationFailed

    async def refuse(db: AsyncSession, req: CalculateRequest) -> CalculateResponse:
        raise ValidationFailed(fields={"crops.0.area": "outside the matrix"})
    monkeypatch.setattr(service.engine, "calculate", refuse)
    fo = await _as(client, office, "field_officer")
    lead = await _lead(client, office, fo)
    r = await _forward(client, fo, lead["id"], office)
    assert r.status_code == 422 and _fields(r) == {"calculation.crops.0.area": "outside the matrix"}, r.text


# ── stages (rules 5 to 7) ────────────────────────────────────────────────────

async def test_stages_move_forward_freely_and_back_only_with_a_remark(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    r = await _record(client, fo, app["id"], "application_in_process", values={"reg_no": "X"})
    assert r.status_code == 422 and "remark" in _fields(r), "same stage needs a remark"
    r = await _record(client, fo, app["id"], "technical_in_process", days_ago=3,
                      values={"tech_received": _ago(3)})
    assert r.status_code == 201, r.text
    assert r.json()["data"]["current_stage"]["code"] == "technical_in_process"
    data = r.json()["data"]
    assert data["ageing"]["days_in_stage"] == _days(data["current_stage"]["since"]) >= 3
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=1)
    assert r.status_code == 422 and "remark" in _fields(r), r.text
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=1, remark="   ")
    assert r.status_code == 422 and "remark" in _fields(r), "a blank remark is no remark"
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=1, remark="Query from GGRC")
    assert r.status_code == 201 and r.json()["data"]["current_stage"]["seq"] == 4, r.text
    entries = (await client.get(f"{APPS}/{app['id']}/stages", headers=fo)).json()["data"]
    assert [e["stage"]["seq"] for e in entries] == [4, 5, 4]
    assert entries[1]["values"] == {"tech_received": _ago(3)} and entries[2]["remark"] == "Query from GGRC"
    assert entries[2]["entered_by"]["id"] == office.shop.ids["field_officer"]


async def test_stage_values_are_checked_field_by_field(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    tomorrow = (today_ist() + dt.timedelta(days=1)).isoformat()
    cases: list[tuple[str, dict[str, Any], str]] = [
        ("technical_in_process", {"tech_received": "29/09/2026"}, "values.tech_received"),
        ("technical_in_process", {"tech_received": tomorrow}, "values.tech_received"),
        ("technical_in_process", {"reg_no": "GGRC-1"}, "values.reg_no"),
        ("farmer_share", {"farmer_share_amt": "100.005"}, "values.farmer_share_amt"),
        ("farmer_share", {"farmer_share_amt": "-1"}, "values.farmer_share_amt"),
        ("farmer_share", {"farmer_share_amt": "NaN"}, "values.farmer_share_amt"),
        ("farmer_share", {"farmer_share_amt": "lots"}, "values.farmer_share_amt"),
        ("farmer_share", {"farmer_share_amt": "1000000000000"}, "values.farmer_share_amt"),
        ("farmer_share", {"farmer_share_amt": "1e30"}, "values.farmer_share_amt"),
        ("tr_pending", {"tpia_name": "x" * 501}, "values.tpia_name"),
        ("no_such_stage", {}, "stage_code"),
    ]
    for stage, values, field in cases:
        r = await _record(client, fo, app["id"], stage, values=values)
        assert r.status_code == 422 and field in _fields(r), (stage, values, r.text)
    r = await client.post(f"{APPS}/{app['id']}/stages", headers={**fo, **_key()},
                          json={"stage_code": "technical_in_process", "occurred_on": tomorrow})
    assert r.status_code == 422 and "occurred_on" in _fields(r), r.text
    entries = (await client.get(f"{APPS}/{app['id']}/stages", headers=fo)).json()["data"]
    assert len(entries) == 1, "no refusal wrote an entry"
    r = await _record(client, fo, app["id"], "farmer_share", values={"farmer_share_amt": "999999999999.99"})
    assert r.status_code == 201, "the widest amount the column holds"
    r = await _record(client, fo, app["id"], "farmer_share", remark="Corrected",
                      values={"farmer_share_amt": "17500.25", "supply": None})
    assert r.status_code == 201, r.text
    entries = (await client.get(f"{APPS}/{app['id']}/stages", headers=fo)).json()["data"]
    assert entries[-1]["values"] == {"farmer_share_amt": "17500.25", "supply": None}


async def test_the_registration_number_and_the_inward_date_follow_the_latest_entry(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    inward = _ago(10)
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=10, remark="Filed",
                      values={"reg_no": " GGRC/24/001 ", "app_inward": inward})
    assert r.status_code == 201, r.text
    data = r.json()["data"]
    assert data["reg_no"] == "GGRC/24/001"
    assert data["ageing"] == {"days_in_stage": _days(data["current_stage"]["since"]), "days_since_inward": _days(inward)}
    await _record(client, fo, app["id"], "technical_in_process", days_ago=9)
    inward = _ago(8)
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=8, remark="Re-filed",
                      values={"reg_no": "GGRC/24/002", "app_inward": inward})
    data = r.json()["data"]
    assert data["reg_no"] == "GGRC/24/002"
    assert data["ageing"] == {"days_in_stage": _days(data["current_stage"]["since"]), "days_since_inward": _days(inward)}
    # a later entry without the number keeps it
    r = await _record(client, fo, app["id"], "technical_in_process", days_ago=2)
    assert r.json()["data"]["reg_no"] == "GGRC/24/002"
    page = (await client.get(APPS, headers=fo, params={"q": "GGRC/24/002"})).json()["data"]
    assert [a["id"] for a in page] == [app["id"]]
    # null or blank withdraws it (review F-2)
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=1, remark="Withdrawn",
                      values={"reg_no": None})
    assert r.json()["data"]["reg_no"] is None, r.text
    await _record(client, fo, app["id"], "application_in_process", days_ago=1, remark="Again",
                  values={"reg_no": "GGRC/24/003"})
    r = await _record(client, fo, app["id"], "application_in_process", days_ago=1, remark="Blank",
                      values={"reg_no": "   "})
    assert r.json()["data"]["reg_no"] is None, r.text
    entries = (await client.get(f"{APPS}/{app['id']}/stages", headers=fo)).json()["data"]
    assert entries[-1]["values"] == {"reg_no": None}, "blank text is stored as cleared"


async def test_the_application_closes_when_every_paid_amount_has_its_date(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    r = await _record(client, fo, app["id"], "payment_received", days_ago=4,
                      values={"pfms_received": _ago(4)})
    assert r.json()["data"]["status"] == "open", "no stage-16 amount: nothing to close on"
    r = await _record(client, fo, app["id"], "fp_cleared_pay_pending", days_ago=3, remark="Cleared",
                      values={"pfms_amt": "50000.00", "state_share_amt": "20000.50",
                              "retention_amt": "0", "total_fp_amt": "70000.50"})
    assert r.status_code == 201 and r.json()["data"]["status"] == "open", r.text
    r = await _record(client, fo, app["id"], "payment_received", days_ago=2,
                      values={"pfms_received": _ago(2)})
    assert r.json()["data"]["status"] == "open", "the state share has no date yet"
    r = await _record(client, fo, app["id"], "payment_received", days_ago=1, remark="State share",
                      values={"state_share_received": _ago(1)})
    data = r.json()["data"]
    # total_fp_amt has no pair and a zero retention needs no date (rule 7)
    assert data["status"] == "full_fp_received" and data["full_fp_received_on"] == _ago(1), data
    r = await _record(client, fo, app["id"], "payment_received", remark="late")
    assert r.status_code == 409 and _code(r) == "status_changed", r.text
    r = await client.post(f"{APPS}/{app['id']}/cancel", headers={**fo, **_key()}, json={"reason": "x"})
    assert r.status_code == 409, r.text
    page = (await client.get(APPS, headers=fo, params={"status": "full_fp_received"})).json()["data"]
    assert app["id"] in [a["id"] for a in page]


async def test_a_cleared_date_holds_closure_back_and_a_cleared_amount_releases_it(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    """Null clears a field: a cleared date is owed again, and an amount cleared
    back out is no longer owed a date."""
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    await _record(client, fo, app["id"], "fp_cleared_pay_pending", days_ago=5,
                  values={"pfms_amt": "10.00", "state_share_amt": "20.00"})
    await _record(client, fo, app["id"], "payment_received", days_ago=4, values={"pfms_received": _ago(4)})
    r = await _record(client, fo, app["id"], "payment_received", days_ago=4, remark="Wrong date",
                      values={"pfms_received": None})
    assert r.json()["data"]["status"] == "open"
    r = await _record(client, fo, app["id"], "payment_received", days_ago=3, remark="State share",
                      values={"state_share_received": _ago(3)})
    assert r.json()["data"]["status"] == "open", "the cleared PFMS date is owed again"
    r = await _record(client, fo, app["id"], "payment_received", days_ago=2, remark="PFMS",
                      values={"pfms_received": _ago(2)})
    assert r.json()["data"]["status"] == "full_fp_received", r.text
    app = await _app(client, office, fo)
    await _record(client, fo, app["id"], "fp_cleared_pay_pending", days_ago=3,
                  values={"pfms_amt": "100.00", "dept_hold_amt": "5.00"})
    r = await _record(client, fo, app["id"], "payment_received", days_ago=2, values={"pfms_received": _ago(2)})
    assert r.json()["data"]["status"] == "open"
    r = await _record(client, fo, app["id"], "fp_cleared_pay_pending", days_ago=1, remark="Hold released",
                      values={"dept_hold_amt": None})
    assert r.json()["data"]["status"] == "full_fp_received", r.text
    assert r.json()["data"]["full_fp_received_on"] == _ago(2)


# ── documents (ADR-041) ──────────────────────────────────────────────────────

class _Counting:
    name = "counting"

    def __init__(self) -> None:
        self.puts: list[str] = []

    def put(self, key: str, data: bytes, content_type: str) -> None:
        self.puts.append(key)

    def presign_get(self, key: str, *, filename: str, disposition: str = "inline") -> tuple[str, dt.datetime]:
        return f"https://files.test/{key}?{disposition}", dt.datetime.now(tz=dt.UTC)


async def _upload(client: httpx.AsyncClient, h: dict[str, str], app_id: str, data: bytes,
                  kind: str = "8a_7_12", name: str = "8a.jpg") -> httpx.Response:
    return await client.post(f"{APPS}/{app_id}/documents", headers={**h, **_key()},
                             files={"file": (name, data, "application/octet-stream")},
                             data={"document_type": kind})


async def test_documents_upload_once_per_file_and_list_on_the_checklist(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        monkeypatch: pytest.MonkeyPatch) -> None:
    storage = _Counting()
    monkeypatch.setattr("api.routers.subsidy_applications.get_storage", lambda *a, **k: storage)
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    r = await _upload(client, fo, app["id"], JPEG, name="8a <scan>.jpg")
    assert r.status_code == 201, r.text
    doc = r.json()["data"]
    assert doc["content_type"] == "image/jpeg" and doc["filename"] == "8a _scan_.jpg"
    r = await _upload(client, fo, app["id"], JPEG, kind="form_16")
    assert r.status_code == 200 and r.json()["data"]["id"] == doc["id"], "the same file is the same document"
    assert len(storage.puts) == 1
    assert (await _upload(client, fo, app["id"], PDF, kind="form_16")).status_code == 201
    r = await _upload(client, fo, app["id"], b"MZ" + b"\x00" * 64)
    assert r.status_code == 422 and _code(r) == "attachment_type", r.text
    r = await _upload(client, fo, app["id"], PDF + b"x", kind="no_such_type")
    assert r.status_code == 422 and "document_type" in _fields(r), r.text
    assert len(storage.puts) == 2

    items = (await client.get(f"{APPS}/{app['id']}/documents", headers=fo)).json()["data"]
    with_files = {i["type"]["code"]: len(i["files"]) for i in items if i["files"]}
    assert with_files == {"8a_7_12": 1, "form_16": 1}
    got = (await client.get(f"{APPS}/{app['id']}", headers=fo)).json()["data"]
    assert got["documents"]["uploaded"] == 2
    r = await client.get(f"{APPS}/{app['id']}/documents/{doc['id']}", headers=fo)
    assert r.status_code == 200 and r.json()["data"]["url"].startswith("https://files.test/subsidy/"), r.text
    r = await client.get(f"{APPS}/{app['id']}/documents/{uuid.uuid4()}", headers=fo)
    assert r.status_code == 404
    timeline = (await client.get(f"{V1}/leads/{app['lead']['id']}/timeline", headers=fo)).json()["data"]
    assert sum(e["kind"] == "subsidy.document_added" for e in timeline) == 2


async def test_a_refused_upload_stores_nothing(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        monkeypatch: pytest.MonkeyPatch) -> None:
    storage = _Counting()
    monkeypatch.setattr("api.routers.subsidy_applications.get_storage", lambda *a, **k: storage)
    fo = await _as(client, office, "field_officer")
    rm = await _as(client, office, "regional_manager")
    app = await _app(client, office, fo)
    r = await _upload(client, rm, app["id"], JPEG)
    assert r.status_code == 403, "view without create or edit adds no files"
    await client.post(f"{APPS}/{app['id']}/cancel", headers={**fo, **_key()}, json={"reason": "Withdrawn"})
    r = await _upload(client, fo, app["id"], JPEG)
    assert r.status_code == 409 and _code(r) == "status_changed", r.text
    assert storage.puts == []
    other = await _app(client, office, fo)
    monkeypatch.setattr("api.routers.subsidy_applications.get_storage",
                        lambda *a, **k: UnconfiguredStorage())
    r = await _upload(client, fo, other["id"], JPEG)
    assert r.status_code == 503 and _code(r) == "storage_unavailable", r.text


async def test_the_fortieth_document_is_the_last(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        sessions: Sessions, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("api.routers.subsidy_applications.get_storage", lambda *a, **k: _Counting())
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    s = sessions()
    await s.execute(text(
        "INSERT INTO subsidy_document (application_id, document_type_id, storage_key, content_type, size_bytes, sha256, uploaded_by) "
        "SELECT CAST(:a AS uuid), (SELECT id FROM subsidy_document_type ORDER BY sort_order LIMIT 1), 'k' || g, 'image/jpeg', 1, "
        "md5(g::text) || md5(:a || g::text), CAST(:u AS uuid) FROM generate_series(1, :n) g"),
        {"a": app["id"], "u": office.shop.ids["field_officer"], "n": service.MAX_DOCUMENTS})
    await s.commit()
    await s.close()
    r = await _upload(client, fo, app["id"], PDF)
    assert r.status_code == 422 and _code(r) == "too_many_documents", r.text


# ── the PIMS sheet (rule 11) ─────────────────────────────────────────────────

async def test_the_pims_sheet_has_a_row_per_line_of_the_calculation(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    from openpyxl import load_workbook

    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    r = await client.get(f"{APPS}/{app['id']}/pims.xlsx", headers=fo)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    rows = [tuple(c.value for c in row) for row in ws.iter_rows()]
    assert rows[0] == service.PIMS_COLUMNS
    body = rows[1:]
    assert body[0][:3] == ("FARMER EDUCATION AND TRAINING", None, "CR 01") and Decimal(str(body[0][8])) == Decimal("500")
    installs = [r for r in body if r[2] == "BQ 01"]
    assert [(r[1], Decimal(str(r[7])), Decimal(str(r[8]))) for r in installs] == [
        ("COTTON", Decimal("1.5"), Decimal("3000")), ("CASTOR", Decimal("0.75"), Decimal("1500"))]
    head = next(r for r in body if r[0] == "Head")
    assert head[3] == "Screen filter" and Decimal(str(head[8])) == Decimal("3500")
    field = next(r for r in body if r[0] == "Field")
    assert (field[1], field[2], field[3]) == ("COTTON", office.item_code, "16 mm lateral")
    assert Decimal(str(field[8])) == Decimal("1250")
    assert len(body) == 5


# ── scope (the `subsidy` ScopeSpec) ──────────────────────────────────────────

async def test_scope_follows_the_role(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    path = f"{APPS}/{app['id']}"
    stranger = await _as(client, office, "stranger")
    # the negative first: territory scope with no territory sees nothing
    assert (await client.get(path, headers=stranger)).status_code == 404
    assert app["id"] not in [a["id"] for a in (await client.get(APPS, headers=stranger)).json()["data"]]
    r = await _record(client, stranger, app["id"], "technical_in_process")
    assert r.status_code == 404, r.text
    for role in ("state_coordinator", "district_manager", "state_manager", "admin_sales", "account_manager"):
        h = await _as(client, office, role)
        assert (await client.get(path, headers=h)).status_code == 200, role
    sc = await _as(client, office, "state_coordinator")
    r = await _record(client, sc, app["id"], "technical_in_process")
    assert r.status_code == 201, "the State Co-ordinator records stages in its state"
    am = await _as(client, office, "account_manager")
    r = await _record(client, am, app["id"], "technical_in_process", remark="x")
    assert r.status_code == 403, "view only"
    r = await client.get(f"{V1}/subsidy-stages", headers=am)
    assert r.status_code == 200 and [s["seq"] for s in r.json()["data"]] == list(range(4, 18))
    closing = next(s for s in r.json()["data"] if s["seq"] == 17)
    assert {f["key"] for f in closing["fields"]} >= {"pfms_received", "state_share_received"}


async def test_a_state_coordinator_forwards_a_lead_in_its_state(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    """The co-ordinator holds no leads.edit; the definer locks the lead itself."""
    fo = await _as(client, office, "field_officer")
    lead = await _lead(client, office, fo)
    sc = await _as(client, office, "state_coordinator")
    r = await _forward(client, sc, lead["id"], office)
    assert r.status_code == 201, r.text
    assert r.json()["data"]["owner"]["id"] == office.shop.ids["field_officer"], "the lead's owner owns it"
    stranger = await _as(client, office, "stranger")
    other = await _lead(client, office, fo)
    r = await _forward(client, stranger, other["id"], office)
    assert r.status_code == 404, r.text


async def test_a_role_without_the_module_is_refused(client: httpx.AsyncClient, office: Office) -> None:
    fo = await _as(client, office, "field_officer")
    r = await client.get(APPS, headers=fo)
    assert r.status_code == 200
    r = await client.get(f"{V1}/subsidy-document-types", headers=fo)
    assert r.status_code == 200 and len(r.json()["data"]) == 20
    qc = await _as(client, office, "qc_manager")
    assert (await client.get(APPS, headers=qc)).status_code == 403


# ── the real engine (skipped without the client's samples) ───────────────────

@needs_samples
async def test_the_drip_sample_forwards_with_the_engines_own_figures(
        client: httpx.AsyncClient, office: Office) -> None:
    from tests.api.test_subsidy_endpoints import _body

    body = _body("drip")
    fo = await _as(client, office, "field_officer")
    calc = await client.post(f"{V1}/subsidy/calculate", headers=fo, json=body)
    assert calc.status_code == 200, calc.text
    result = calc.json()["data"]
    category = next(c["code"] for c in result["crops"][0]["categories"]
                    if all(any(x["code"] == c["code"] and x["applicable"] for x in crop["categories"])
                           for crop in result["crops"]))
    lead = await _lead(client, office, fo)
    r = await client.post(APPS, headers={**fo, **_key()},
                          json={"lead_id": lead["id"], "category_code": category, "calculation": body})
    assert r.status_code == 201, r.text
    app = r.json()["data"]
    rows = [next(x for x in crop["categories"] if x["code"] == category) for crop in result["crops"]]
    assert Decimal(app["figures"]["subsidy"]) == sum(Decimal(x["subsidy"]) for x in rows)
    assert Decimal(app["figures"]["farmer_share"]) == sum(Decimal(x["farmer_share"]) for x in rows)
    assert app["figures"]["total_cost"] == result["total"]["blocks"]["total_incl_gst"]
    stored = (await client.get(f"{APPS}/{app['id']}/calculation", headers=fo)).json()["data"]
    assert stored == result


# ── code review (Fable, on the build) ────────────────────────────────────────

async def test_a_dealer_is_refused_on_every_subsidy_route(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        sessions: Sessions) -> None:
    """F-5: no dealer role holds subsidy."""
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    dealer, mobile = await complaints._dealer(client, office.shop, sessions)
    try:
        for method, path in (("GET", APPS), ("GET", f"{APPS}/{app['id']}"),
                             ("GET", f"{APPS}/{app['id']}/stages"), ("GET", f"{APPS}/{app['id']}/pims.xlsx"),
                             ("GET", f"{APPS}/{app['id']}/documents"), ("GET", f"{V1}/subsidy-stages"),
                             ("GET", f"{V1}/subsidy-document-types")):
            r = await client.request(method, path, headers=dealer)
            assert r.status_code == 403, (path, r.status_code)
        r = await _record(client, dealer, app["id"], "technical_in_process")
        assert r.status_code == 403
        r = await _forward(client, dealer, app["lead"]["id"], office)
        assert r.status_code == 403
    finally:
        await complaints._forget(sessions, mobile)


def _caller(o: Office) -> Caller:
    return Caller(o.shop.ids["field_officer"], o.shop.office, None, scopes={"subsidy": "own"})


async def test_two_forwards_of_one_lead_leave_one_application(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        sessions: Sessions) -> None:
    """F-3: the second waits on the first's lead lock, then finds the application."""
    fo = await _as(client, office, "field_officer")
    lead = await _lead(client, office, fo)
    body = sch.ApplicationCreate.model_validate(
        {"lead_id": lead["id"], "category_code": "small_farmer", "calculation": _calc(office)})

    async def work(s: AsyncSession) -> Any:
        return await service.create(s, _caller(office), body)
    me = office.shop.ids["field_officer"]
    got, waited = await conc._race(sessions, (me, work), (me, work), on="lead")
    assert waited, "the second forward did not queue on the lead's row lock"
    assert sorted(conc._outcome(g) for g in got) == ["already_forwarded", "ok"], got


async def test_two_closing_entries_close_once(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        sessions: Sessions) -> None:
    """F-3: the second waits on the application's lock and finds it closed.
    Without the lock both entries land and the application closes twice."""
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    await _record(client, fo, app["id"], "fp_cleared_pay_pending", days_ago=2, values={"pfms_amt": "10.00"})
    body = sch.StageRecord.model_validate({"stage_code": "payment_received", "occurred_on": _ago(1),
                                           "values": {"pfms_received": _ago(1)}, "remark": "Paid"})

    async def work(s: AsyncSession) -> Any:
        return await service.record(s, app["id"], body)
    me = office.shop.ids["field_officer"]
    got, waited = await conc._race(sessions, (me, work), (me, work))
    assert waited, "the second entry did not wait: the race was not a race"
    assert sorted(conc._outcome(g) for g in got) == ["ok", "status_changed"], got
    timeline = (await client.get(f"{V1}/leads/{app['lead']['id']}/timeline", headers=fo)).json()["data"]
    assert sum(e["kind"] == "subsidy.closed" for e in timeline) == 1


async def test_two_uploads_at_the_cap_leave_forty(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        sessions: Sessions) -> None:
    """F-3: at 39 files the second upload waits on the first's lock and counts 40."""
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    s = sessions()
    await s.execute(text(
        "INSERT INTO subsidy_document (application_id, document_type_id, storage_key, content_type, size_bytes, sha256, uploaded_by) "
        "SELECT CAST(:a AS uuid), (SELECT id FROM subsidy_document_type ORDER BY sort_order LIMIT 1), 'k' || g, 'image/jpeg', 1, "
        "md5(g::text) || md5(:a || g::text), CAST(:u AS uuid) FROM generate_series(1, :n) g"),
        {"a": app["id"], "u": office.shop.ids["field_officer"], "n": service.MAX_DOCUMENTS - 1})
    await s.commit()
    await s.close()

    def upload(data: bytes) -> conc.Work:
        async def work(db: AsyncSession) -> Any:
            return await service.add_document(db, _caller(office), app["id"], document_type="form_16",
                                              filename="f.pdf", data=data, storage=_Counting())
        return work
    me = office.shop.ids["field_officer"]
    got, waited = await conc._race(sessions, (me, upload(PDF)), (me, upload(PDF + b"other")))
    assert waited, "the second upload did not wait: the race was not a race"
    assert sorted(conc._outcome(g) for g in got) == ["ok", "too_many_documents"], got


# ── OpenCodeReview (PR 33) ───────────────────────────────────────────────────

async def test_stage_values_are_strings_only(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    """A JSON number would be a float for money; lax mode read `true` as 1."""
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    for value in (17500.25, 100, True):
        r = await _record(client, fo, app["id"], "farmer_share", values={"farmer_share_amt": value})
        assert r.status_code == 422, (value, r.text)


async def test_a_blank_date_or_amount_clears_the_field(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    """An emptied date input sends an empty string."""
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    r = await _record(client, fo, app["id"], "farmer_share", values={"supply": "", "farmer_share_amt": " "})
    assert r.status_code == 201, r.text
    entries = (await client.get(f"{APPS}/{app['id']}/stages", headers=fo)).json()["data"]
    assert entries[-1]["values"] == {"supply": None, "farmer_share_amt": None}


async def test_values_on_a_hidden_application_or_an_unknown_stage_name_the_right_field(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest]) -> None:
    fo = await _as(client, office, "field_officer")
    app = await _app(client, office, fo)
    stranger = await _as(client, office, "stranger")
    r = await _record(client, stranger, app["id"], "technical_in_process", values={"tech_received": _ago(1)})
    assert r.status_code == 404, r.text
    r = await _record(client, fo, app["id"], "no_such_stage", values={"tech_received": _ago(1)})
    assert r.status_code == 422 and _fields(r) == {"stage_code": "not a stage of this scheme"}, r.text


async def test_a_reused_application_number_is_a_fault_not_already_forwarded(
        client: httpx.AsyncClient, office: Office, fake_engine: list[CalculateRequest],
        sessions: Sessions) -> None:
    """Only the one-live-application index means `already_forwarded`."""
    from sqlalchemy.exc import DBAPIError

    fo = await _as(client, office, "field_officer")
    await _app(client, office, fo)
    s = sessions()
    await s.execute(text("UPDATE subsidy_app_counter SET last_value = 0 WHERE state_code = :c"),
                    {"c": office.shop.code})
    await s.commit()
    await s.close()
    lead = await _lead(client, office, fo)
    try:
        r = await _forward(client, fo, lead["id"], office)
    except DBAPIError as exc:  # the test client re-raises what the app does not map
        assert getattr(exc.orig, "sqlstate", None) == "23505", exc
    else:
        assert r.status_code == 500, r.text


def test_a_definer_message_loses_any_class_prefix() -> None:
    """SQLAlchemy 2.0's asyncpg adapter prefixes the exception class."""
    from sqlalchemy.exc import DBAPIError

    for raw in ("tech_received", "<class 'asyncpg.exceptions.RaiseError'>: tech_received"):
        assert service._message(DBAPIError("SELECT 1", {}, Exception(raw))) == "tech_received"


def test_the_upload_route_documents_its_refusals() -> None:
    from api.main import app as api

    op = api.openapi()["paths"]["/api/v1/subsidy-applications/{app_id}/documents"]["post"]
    assert {"403", "409", "413", "422", "503"} <= set(op["responses"])
    for code in ("too_many_documents", "attachment_type", "not on the checklist"):
        assert code in op["description"], code
    link = api.openapi()["paths"]["/api/v1/subsidy-applications/{app_id}/documents/{doc_id}"]["get"]
    assert "503" in link["responses"]


def test_the_violated_constraint_is_read_from_the_driver_error_or_its_cause() -> None:
    """asyncpg's error sits under the adapter's as `__cause__`; psycopg's under `diag`."""
    from types import SimpleNamespace

    from sqlalchemy.exc import DBAPIError

    cause = Exception("duplicate")
    cause.constraint_name = "uq_subsidy_application_live_lead"  # type: ignore[attr-defined]
    adapted = Exception("wrapped")
    adapted.__cause__ = cause
    assert service._constraint(DBAPIError("x", {}, adapted)) == "uq_subsidy_application_live_lead"
    psycopg_like = Exception("dup")
    psycopg_like.diag = SimpleNamespace(constraint_name="uq_subsidy_document_file")  # type: ignore[attr-defined]
    assert service._constraint(DBAPIError("x", {}, psycopg_like)) == "uq_subsidy_document_file"
    assert service._constraint(DBAPIError("x", {}, Exception("other"))) is None
