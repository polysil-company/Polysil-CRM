# ruff: noqa: E501  (sheet data)

"""The import template endpoint (FS-068): a file per kind, the upload's
permission, not caught by /imports/{batch_id}, refused when uploaded blank, and
matched by name when filled, even after a wrong mapping was saved.

Built on FS-062's `imported` fixture, which removes batches and lists afterwards."""

from __future__ import annotations

import io
from collections.abc import Callable

import httpx
import openpyxl
import pytest
from sqlalchemy.ext.asyncio import AsyncSession

from api.domain.imports.fields import FIELDS
from api.errors import ForbiddenError
from api.services import imports as service
from tests.api import test_imports as base
from tests.api import test_order_endpoints as endpoints
from tests.api.conftest import V1
from tests.api.test_order_partners import _as_dealer

pytestmark = pytest.mark.db

shop = base.shop
imported = base.imported
__all__ = ["imported", "shop"]

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


async def _template(client: httpx.AsyncClient, h: dict[str, str], kind: str) -> httpx.Response:
    return await client.get(f"{V1}/imports/template", headers=h, params={"kind": kind})


@pytest.mark.parametrize(("kind", "name"), [("partners", "partner-import-template.xlsx"),
                                            ("price_list_items", "price-list-template.xlsx"),
                                            ("lines", "order-lines-template.xlsx")])
async def test_each_kind_downloads_its_file(client: httpx.AsyncClient, imported: base.Fx,
                                            kind: str, name: str) -> None:
    h = await endpoints._as(client, imported.shop, "admin_sales")
    r = await _template(client, h, kind)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == f'attachment; filename="{name}"'
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).worksheets[0]
    assert [c.value for c in ws[1]] == [f.label for f in FIELDS[kind]]


async def test_a_missing_or_unknown_kind(client: httpx.AsyncClient, imported: base.Fx) -> None:
    h = await endpoints._as(client, imported.shop, "admin_sales")
    assert (await _template(client, h, "staff")).status_code == 422
    r = await client.get(f"{V1}/imports/template", headers=h)
    assert r.status_code == 422
    assert "kind" in r.text and "batch_id" not in r.text, "caught by /imports/{batch_id}"


async def test_not_signed_in(client: httpx.AsyncClient) -> None:
    assert (await client.get(f"{V1}/imports/template",
                             params={"kind": "lines"})).status_code == 401


async def test_a_dealer_gets_the_lines_template_only(
        imported: base.Fx, sessions: Callable[[], AsyncSession]) -> None:
    """Rule 4: the upload's permission. Partner and price imports are staff only."""
    s, caller = await _as_dealer(sessions, imported.shop)
    try:
        for kind in ("partners", "price_list_items"):
            with pytest.raises(ForbiddenError):
                await service.template(s, caller, kind)
        assert (await service.template(s, caller, "lines")).startswith(b"PK")
    finally:
        await s.rollback()
        await s.close()


async def test_the_blank_template_is_refused_at_upload(
        client: httpx.AsyncClient, imported: base.Fx) -> None:
    """Rule 7: before, it uploaded as 0 rows and checked as 0 of 0."""
    h = await endpoints._as(client, imported.shop, "admin_sales")
    blank = (await _template(client, h, "lines")).content
    r = await base._upload(client, h, "lines", blank, name="order-lines-template.xlsx")
    assert r.status_code == 422, r.text
    body = r.json()
    assert body["error"]["code"] == "import_unreadable", body
    assert "no rows under them" in body["error"]["fields"]["file"]


async def test_a_chosen_header_row_with_no_rows_is_reported_on_header_row(
        client: httpx.AsyncClient, imported: base.Fx) -> None:
    """Code review F-2: with both sheet and header_row sent, the error names the
    header row, not the sheet."""
    h = await endpoints._as(client, imported.shop, "admin_sales")
    blank = (await _template(client, h, "lines")).content
    r = await base._upload(client, h, "lines", blank, name="t.xlsx", sheet="Order lines",
                           header_row="1")
    assert r.status_code == 422, r.text
    assert r.json()["error"]["fields"] == {"header_row": "Row 1 has no rows under it."}


async def test_a_filled_template_is_matched_by_name_after_a_wrong_mapping(
        client: httpx.AsyncClient, imported: base.Fx) -> None:
    """Rule 6 and EC-1: the 16 Oct file, names under "Item Code", saved as the
    description. The template's Item code column must not inherit it."""
    h = await endpoints._as(client, imported.shop, "admin_sales")
    list_id = await base._draft_list(client, h)
    wrong = (await base._upload(client, h, "price_list_items", base._csv([
        ["Item Code", "Rate"], [imported.product_description, "5"],
    ]), name="r.csv", target_id=list_id)).json()["data"]
    assert (await base._check(client, h, wrong["id"], {"description": 0, "rate": 1})).status_code == 200
    assert (await base._commit(client, h, wrong["id"])).status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO((await _template(client, h, "price_list_items")).content))
    wb.worksheets[0].append(["EXAMPLE-1", imported.product_description, 7])
    out = io.BytesIO()
    wb.save(out)
    batch = (await base._upload(client, h, "price_list_items", out.getvalue(),
                                name="price-list-template.xlsx", target_id=list_id)).json()["data"]
    assert batch["header_row"] == 1
    assert batch["suggested_mapping"] == {"item_code": 0, "description": 1, "rate": 2}
