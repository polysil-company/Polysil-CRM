"""FS-024 over the API: each report's shape and counts on known data, the null
figures, the staff-only 403s, Excel, and the 360."""

# ruff: noqa: E501  (request bodies inline)

from __future__ import annotations

import io
import uuid
from collections.abc import Callable
from decimal import Decimal

import httpx
import pytest
from argon2 import PasswordHasher
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from api.services import exports, report_exports
from tests.api import test_complaints as complaints_t
from tests.api import test_notifications as notif
from tests.api import test_order_endpoints as endpoints
from tests.api import test_payments as payments_t
from tests.api.conftest import PASSWORD, V1, _key
from tests.api.test_payments import (
    _forget_payments,  # noqa: F401  (autouse: payment rows before the shop)
)
from tests.api.test_settings_amend_reopen import restore  # noqa: F401  (fixture)

pytestmark = pytest.mark.db

shop = endpoints.shop
Shop = endpoints.Shop
Sessions = Callable[[], AsyncSession]


async def _lost(client: httpx.AsyncClient, shop: Shop, h: dict[str, str]) -> dict:
    lead = await notif._lead(client, shop, h)
    reasons = [x for x in (await client.get(f"{V1}/lookups/lost-reasons", headers=h)).json()["data"] if x["is_active"]]
    r = await client.post(f"{V1}/leads/{lead['id']}/transition", headers={**h, **_key()},
                          json={"to_stage": "lost", "lost_reason_id": reasons[0]["id"]})
    assert r.status_code == 200, r.text
    return lead


async def test_conversion_counts_my_leads_and_lost_leads_carry_their_stage(client: httpx.AsyncClient, shop: Shop) -> None:
    fo = await endpoints._as(client, shop, "field_officer")
    for _ in range(2):
        await notif._lead(client, shop, fo)
    await _lost(client, shop, fo)
    r = await client.get(f"{V1}/reports/lead-conversion", headers=fo, params={"group_by": "owner", "territory_id": shop.district})
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    mine = next(x for x in data["rows"] if x["key"]["id"] == shop.ids["field_officer"])
    assert mine["leads"] == 3 and mine["lost"] == 1 and mine["open"] == 2 and mine["conversion_pct"] == "0.0"
    assert data["totals"]["leads"] == 3, "a field officer's report is about their own leads"
    lost = (await client.get(f"{V1}/reports/lost-leads", headers=fo, params={"territory_id": shop.district})).json()["data"]
    assert lost["totals"]["lost"] == 1 and lost["rows"][0]["by_stage"] == {"new": 1}


async def test_a_figure_without_its_module_is_null_never_zero(client: httpx.AsyncClient, shop: Shop) -> None:
    """Review B-1: Accounts holds no tasks or tracking; those figures are null."""
    await notif._lead(client, shop, await endpoints._as(client, shop, "field_officer"))
    accounts = await endpoints._as(client, shop, "account_manager")
    r = await client.get(f"{V1}/reports/salesperson-performance", headers=accounts, params={"territory_id": shop.district})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["rows"], "a row to check: the officer's lead (code review)"
    for row in r.json()["data"]["rows"]:
        assert row["tasks_done"] is None and row["visits"] is None
        assert row["orders"] is not None
    x = await client.get(f"{V1}/reports/salesperson-performance", headers=accounts,
                         params={"territory_id": shop.district, "format": "xlsx"})
    assert x.status_code == 200, x.text
    sheet = [tuple(c.value for c in row) for row in load_workbook(io.BytesIO(x.content)).active.iter_rows()]
    visits, tasks = sheet[0].index("Visits"), sheet[0].index("Tasks done")
    assert len(sheet) > 2, "data rows besides the header and the Total line"
    assert all(row[visits] is None and row[tasks] is None for row in sheet[1:]), \
        "a figure the caller may not see is a blank cell in Excel too, never 0"
    # the positive control: a district manager's task scope reaches the officer
    dm = await endpoints._as(client, shop, "district_manager")
    mine = (await client.get(f"{V1}/reports/salesperson-performance", headers=dm,
                             params={"territory_id": shop.district})).json()["data"]["rows"]
    officer = next(x for x in mine if x["user"]["id"] == shop.ids["field_officer"])
    assert officer["tasks_done"] == 0 and officer["visits"] == 0, "a real 0 where the scope reaches"
    r = await client.get(f"{V1}/reports/follow-ups", headers=accounts)
    assert r.status_code == 200, "Accounts holds its own tasks"


async def test_staff_reports_refuse_a_dealer_and_excel_downloads(client: httpx.AsyncClient, shop: Shop,
                                                                     sessions: Sessions) -> None:
    dealer, mobile = await complaints_t._dealer(client, shop, sessions)
    try:
        for name in ("salesperson-performance", "follow-ups", "territory-performance"):
            assert (await client.get(f"{V1}/reports/{name}", headers=dealer)).status_code == 403, name
    finally:
        await complaints_t._forget(sessions, mobile)
    dm = await endpoints._as(client, shop, "district_manager")
    r = await client.get(f"{V1}/reports/lead-conversion", headers=dm, params={"format": "xlsx"})
    assert r.status_code == 200 and r.headers["content-type"].startswith(exports.XLSX_TYPE), r.text
    assert 'filename="report-lead-conversion-' in r.headers["content-disposition"]
    sheet = [tuple(c.value for c in row) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows()]
    assert sheet[0][:2] == ("Source", "Leads") and sheet[-1][0] == "Total", "GAP-220: the rows, then a Total line"
    r = await client.get(f"{V1}/reports/lead-conversion", headers=dm, params={"from": "2026-10-05", "to": "2026-10-01"})
    assert r.status_code == 422


def _has_path(item: dict, path: str) -> bool:
    for part in path.split("."):
        if not isinstance(item, dict) or part not in item:
            return False
        item = item[part]
    return True


async def test_every_report_answers_for_a_manager(client: httpx.AsyncClient, shop: Shop) -> None:
    fo = await endpoints._as(client, shop, "field_officer")
    await notif._lead(client, shop, fo)               # rows for the lead and person reports
    await _lost(client, shop, fo)                     # a row for lost leads
    await endpoints._submit(client, fo, (await endpoints._create(
        client, fo, endpoints._direct(shop, partner_id=shop.partner)))["id"])  # the dealer report
    dm = await endpoints._as(client, shop, "district_manager")
    for name, params in (("lead-conversion", {}), ("salesperson-performance", {}), ("lost-leads", {}),
                         ("follow-ups", {}), ("dealer-performance", {}),
                         ("territory-performance", {"level": "district"}), ("complaints", {})):
        r = await client.get(f"{V1}/reports/{name}", headers=dm, params=params)
        assert r.status_code == 200, (name, r.text)
        data = r.json()["data"]
        assert {"rows", "totals", "truncated", "filters"} <= set(data), name
        # code review: every Excel column names a key the service really returns, so
        # a renamed key cannot ship a silently blank column
        cols = report_exports.columns(name, data)
        for row in data["rows"][:1]:
            for col in cols:
                # by_stage holds only the stages that occur; its parent must exist
                path = "by_stage" if col.path.startswith("by_stage.") else col.path
                assert _has_path(row, path), (name, col.path)
        assert any(_has_path(data["totals"], k) for k in data["totals"]), name
        x = await client.get(f"{V1}/reports/{name}", headers=dm, params={**params, "format": "xlsx"})
        assert x.status_code == 200, (name, x.text)
        sheet = [tuple(c.value for c in row) for row in load_workbook(io.BytesIO(x.content)).active.iter_rows()]
        assert sheet[0][0] == cols[0].header and "Total" in [row[0] for row in sheet], name
        assert len(sheet) == len(data["rows"]) + 2 + (name == "complaints") + bool(data["truncated"]), name


async def test_the_dealer_report_counts_orders_and_payments(client: httpx.AsyncClient, shop: Shop) -> None:
    fo = await endpoints._as(client, shop, "field_officer")
    order = await endpoints._approve_all(client, shop, await endpoints._submit(
        client, fo, (await endpoints._create(client, fo, endpoints._direct(shop, partner_id=shop.partner)))["id"]))
    admin = await endpoints._as(client, shop, "admin_sales")
    data = (await client.get(f"{V1}/reports/dealer-performance", headers=admin, params={"territory_id": shop.district})).json()["data"]
    row = next(x for x in data["rows"] if x["partner"]["id"] == shop.partner)
    assert row["orders"] == 1 and row["order_value"] == order["totals"]["total"]
    assert row["received"] == "0.00" and row["balance"] == order["totals"]["total"]


async def test_the_360_shows_tiles_and_same_mobile_leads(client: httpx.AsyncClient, shop: Shop) -> None:
    fo = await endpoints._as(client, shop, "field_officer")
    mobile = "97" + f"{uuid.uuid4().int % 10**8:08d}"
    ids = []
    for _ in range(2):
        r = await client.post(f"{V1}/leads", headers={**fo, **_key()}, json={
            "farmer_name": "Same Phone", "mobile": mobile, "territory_id": shop.district,
            "inquiry_type": "commercial", "mis_system": "drip"})
        assert r.status_code in (200, 201), r.text
        ids.append(r.json()["data"]["id"])
    r = await client.get(f"{V1}/leads/{ids[0]}/360", headers=fo)
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    assert [x["id"] for x in data["related_leads"]] == [ids[1]] and data["related_leads"][0]["relation"] == "same_mobile"
    assert data["summary"]["received"] is None, "a field officer holds no payments.view"
    assert data["summary"]["open_tasks"] == 0


async def _second_office(client: httpx.AsyncClient, shop: Shop, sessions: Sessions) -> dict[str, str]:
    """A sibling office in the same district with its own field officer. The shop's
    teardown removes both: the person through `ids`, the office by district."""
    s = sessions()
    try:
        office = str((await s.execute(text(
            "INSERT INTO org_unit (name, role_level, territory_id) VALUES (:n, 2, CAST(:t AS uuid)) RETURNING id"),
            {"n": f"report_office_b_{uuid.uuid4().hex[:6]}", "t": shop.district})).scalar_one())
        email = f"report_fo_b_{uuid.uuid4().hex[:8]}@polysil.in"
        fo_b = str((await s.execute(text(
            "INSERT INTO app_user (user_type, email, password_hash, full_name, role_id, org_unit_id) "
            "SELECT 'staff', :e, :p, 'Other Office Officer', r.id, CAST(:o AS uuid) FROM role r "
            "WHERE r.code = 'field_officer' RETURNING id"),
            {"e": email, "p": PasswordHasher().hash(PASSWORD), "o": office})).scalar_one())
        await s.commit()
    finally:
        await s.close()
    shop.ids["field_officer_b"], shop.users["field_officer_b"] = fo_b, email
    return await endpoints._as(client, shop, "field_officer_b")


async def test_a_manager_counts_their_office_only_and_the_dealer_balance_agrees(
        client: httpx.AsyncClient, shop: Shop, sessions: Sessions, restore: None) -> None:  # noqa: F811
    """Review 1 and 2: a sibling office's leads, orders and receipts stay out of a
    District Manager's figures, and received is over the same orders as payable."""
    fo_b = await _second_office(client, shop, sessions)
    fo = await endpoints._as(client, shop, "field_officer")
    await notif._lead(client, shop, fo)
    await notif._lead(client, shop, fo_b)
    mine = await endpoints._submit(client, fo, (await endpoints._create(
        client, fo, endpoints._direct(shop, partner_id=shop.partner)))["id"])
    theirs = await endpoints._submit(client, fo_b, (await endpoints._create(
        client, fo_b, endpoints._direct(shop, partner_id=shop.partner)))["id"])
    # the sibling office has no approvers of its own; this test is about scope,
    # so it counts sales from submission (FS-026), restored by `restore`
    s = sessions()
    await s.execute(text("UPDATE app_setting SET value = '\"submission\"' WHERE key = 'sale_counted_at'"))
    await s.commit()
    await s.close()
    accounts = await endpoints._as(client, shop, "account_manager")
    r = await payments_t._pay(client, accounts, {"partner_id": shop.partner, "amount": theirs["totals"]["total"],
                                                 "allocations": [{"sales_order_id": theirs["id"],
                                                                  "amount": theirs["totals"]["total"]}]})
    assert r.status_code == 201, r.text
    p = {"territory_id": shop.district}
    dm = await endpoints._as(client, shop, "district_manager")
    admin = await endpoints._as(client, shop, "admin_sales")
    for h, leads in ((dm, 1), (admin, 2)):
        data = (await client.get(f"{V1}/reports/lead-conversion", headers=h, params=p)).json()["data"]
        assert data["totals"]["leads"] == leads, data["totals"]
    row = next(x for x in (await client.get(f"{V1}/reports/dealer-performance", headers=dm, params=p)).json()["data"]["rows"]
               if x["partner"]["id"] == shop.partner)
    assert row["orders"] == 1 and row["received"] == "0.00" and row["balance"] == mine["totals"]["total"], row
    row = next(x for x in (await client.get(f"{V1}/reports/dealer-performance", headers=admin, params=p)).json()["data"]["rows"]
               if x["partner"]["id"] == shop.partner)
    assert row["orders"] == 2 and row["received"] == theirs["totals"]["total"], row
    assert Decimal(row["balance"]) == Decimal(mine["totals"]["total"]), row


async def test_territory_performance_keeps_leads_above_the_level(client: httpx.AsyncClient, shop: Shop) -> None:
    """Review 4: the shop's leads sit at the district, so by taluka they are one row."""
    fo = await endpoints._as(client, shop, "field_officer")
    await notif._lead(client, shop, fo)
    dm = await endpoints._as(client, shop, "district_manager")
    data = (await client.get(f"{V1}/reports/territory-performance", headers=dm,
                             params={"level": "taluka", "territory_id": shop.district})).json()["data"]
    assert data["totals"]["leads"] == 1
    assert data["rows"][0]["territory"] == {"id": None, "name": "Above this level"}
