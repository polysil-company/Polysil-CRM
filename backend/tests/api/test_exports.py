"""FS-030: every export holds exactly the rows its list would show, across every
page, and nothing outside the caller's scope. The negative cases are the point:
an export is the easiest data leak in the system (Build-Plan 9.3)."""

# ruff: noqa: E501  (embedded SQL)

from __future__ import annotations

import io
import uuid
from collections.abc import Callable
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from api.services import exports
from tests.api import test_complaints as complaints
from tests.api import test_order_endpoints as endpoints
from tests.api import test_subsidy_applications as subsidy
from tests.api.conftest import V1

shop = endpoints.shop
Shop = endpoints.Shop
office = subsidy.office
fake_engine = subsidy.fake_engine
Sessions = Callable[[], AsyncSession]

pytestmark = pytest.mark.db


def _rows(r: httpx.Response) -> list[dict[str, Any]]:
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == exports.XLSX_TYPE
    assert r.headers["content-disposition"].startswith('attachment; filename="')
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    return [dict(zip(header, (c.value for c in row), strict=True)) for row in ws.iter_rows(min_row=2)]


async def _leads(sessions: Sessions, shop: Shop, tag: str, owners: list[str]) -> list[str]:
    s = sessions()
    ids = []
    try:
        for i, owner in enumerate(owners):
            ids.append(str((await s.execute(text(
                "INSERT INTO lead (inquiry_no, inquiry_type, mis_system_id, lead_source_id, farmer_name, "
                "mobile, territory_id, owner_user_id, owner_org_unit_id, created_by, estimated_value) "
                "VALUES (:no, 'commercial', (SELECT id FROM mis_system WHERE code = 'drip'), "
                "(SELECT id FROM lead_source WHERE code = 'employee'), :name, :mob, CAST(:t AS uuid), "
                "CAST(:o AS uuid), CAST(:ou AS uuid), CAST(:o AS uuid), 125000.50) RETURNING id"),
                {"no": f"EXP-{tag}-{i}", "name": f"Export Farmer {tag} {i}",
                 "mob": "+9197" + f"{uuid.uuid4().int % 10**8:08d}", "t": shop.district,
                 "o": owner, "ou": shop.office})).scalar_one()))
        await s.commit()
    finally:
        await s.close()
    return ids


async def test_lead_export_equals_the_list_across_pages_and_keeps_scope(
        client: httpx.AsyncClient, shop: Shop, sessions: Sessions,
        monkeypatch: pytest.MonkeyPatch) -> None:
    tag = uuid.uuid4().hex[:8]
    fo, dm = shop.ids["field_officer"], shop.ids["district_manager"]
    await _leads(sessions, shop, tag, [fo, fo, dm])
    monkeypatch.setattr(exports, "PAGE", 1)     # three pages, so the drain is exercised

    admin = await endpoints._as(client, shop, "admin_sales")
    rows = _rows(await client.get(f"{V1}/leads/export", headers=admin, params={"q": f"Export Farmer {tag}"}))
    listed = (await client.get(f"{V1}/leads", headers=admin, params={"q": f"Export Farmer {tag}"})).json()["data"]
    assert sorted(r["Inquiry no"] for r in rows) == sorted(x["inquiry_no"] for x in listed)
    assert len(rows) == 3
    assert rows[0]["Estimated value"] == 125000.5, "money is a number"
    assert rows[0]["Territory"].startswith("ord_district_")

    # the negative: a field officer downloads their own two, never the manager's
    h = await endpoints._as(client, shop, "field_officer")
    mine = _rows(await client.get(f"{V1}/leads/export", headers=h, params={"q": f"Export Farmer {tag}"}))
    assert sorted(r["Inquiry no"] for r in mine) == [f"EXP-{tag}-0", f"EXP-{tag}-1"]
    assert f"EXP-{tag}-2" not in {r["Inquiry no"] for r in mine}

    # a filter narrows the export the way it narrows the list
    one = _rows(await client.get(f"{V1}/leads/export", headers=admin,
                                 params={"q": f"Export Farmer {tag}", "owner_user_id": dm}))
    assert [r["Inquiry no"] for r in one] == [f"EXP-{tag}-2"]


async def test_more_than_the_cap_is_refused_and_exactly_the_cap_is_not(
        client: httpx.AsyncClient, shop: Shop, sessions: Sessions,
        monkeypatch: pytest.MonkeyPatch) -> None:
    tag = uuid.uuid4().hex[:8]
    fo = shop.ids["field_officer"]
    await _leads(sessions, shop, tag, [fo, fo, fo])
    h = await endpoints._as(client, shop, "field_officer")
    monkeypatch.setattr(exports, "MAX_ROWS", 3)
    assert len(_rows(await client.get(f"{V1}/leads/export", headers=h,
                                      params={"q": f"Export Farmer {tag}"}))) == 3
    monkeypatch.setattr(exports, "MAX_ROWS", 2)
    r = await client.get(f"{V1}/leads/export", headers=h, params={"q": f"Export Farmer {tag}"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "export_too_large"


async def test_without_the_view_permission_the_export_is_403(
        client: httpx.AsyncClient, shop: Shop) -> None:
    h = await endpoints._as(client, shop, "field_officer")     # no users.view (RBAC 6.1)
    r = await client.get(f"{V1}/users/export", headers=h)
    assert r.status_code == 403
    r = await client.get(f"{V1}/leads/export")
    assert r.status_code == 401


async def test_every_export_answers_with_a_workbook(client: httpx.AsyncClient, shop: Shop) -> None:
    h = await endpoints._as(client, shop, "admin_sales")
    for path in ("quotations", "orders", "complaints", "tasks", "subsidy-applications",
                 "partners", "users"):
        r = await client.get(f"{V1}/{path}/export", headers=h, params={"q": "no such row zz"}
                             if path not in ("tasks",) else {"task_type": "call"})
        _rows(r)


async def test_the_complaint_queue_filter_is_refused_not_cut(
        client: httpx.AsyncClient, shop: Shop) -> None:
    h = await endpoints._as(client, shop, "district_manager")
    r = await client.get(f"{V1}/complaints/export", headers=h, params={"awaiting": "me"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "filter_not_exportable"


async def test_complaint_export_keeps_a_field_officer_to_their_own(
        client: httpx.AsyncClient, shop: Shop) -> None:
    fo = await endpoints._as(client, shop, "field_officer")
    dm = await endpoints._as(client, shop, "district_manager")
    tag = uuid.uuid4().hex[:6]
    mine = await complaints._create(client, shop, fo, contact_name=f"Export Mine {tag}")
    theirs = await complaints._create(client, shop, dm, contact_name=f"Export Theirs {tag}")
    rows = _rows(await client.get(f"{V1}/complaints/export", headers=fo, params={"q": "Export"}))
    contacts = {r["Contact"] for r in rows}
    assert f"Export Mine {tag}" in contacts
    assert f"Export Theirs {tag}" not in contacts
    every = _rows(await client.get(f"{V1}/complaints/export", headers=dm, params={"q": "Export"}))
    assert {f"Export Mine {tag}", f"Export Theirs {tag}"} <= {r["Contact"] for r in every}
    assert mine["id"] != theirs["id"]


async def test_task_export_keeps_a_field_officer_to_their_own(
        client: httpx.AsyncClient, shop: Shop, sessions: Sessions) -> None:
    tag = uuid.uuid4().hex[:8]
    s = sessions()
    try:
        for who in ("field_officer", "district_manager"):
            await s.execute(text(
                "INSERT INTO task (title, task_type, status, due_at, assigned_to, assigned_by, owner_org_unit_id) "
                "VALUES (:t, 'call', 'open', now() + interval '1 day', CAST(:u AS uuid), CAST(:b AS uuid), "
                "CAST(:ou AS uuid))"),
                {"t": f"Export task {tag} {who}", "u": shop.ids[who],
                 "b": shop.ids["district_manager"], "ou": shop.office})
        await s.commit()
    finally:
        await s.close()
    fo = await endpoints._as(client, shop, "field_officer")
    titles = {r["Title"] for r in _rows(await client.get(f"{V1}/tasks/export", headers=fo))}
    assert f"Export task {tag} field_officer" in titles
    assert f"Export task {tag} district_manager" not in titles


async def test_partner_export_hides_credit_terms_without_partners_edit(
        client: httpx.AsyncClient, shop: Shop, sessions: Sessions) -> None:
    s = sessions()
    await s.execute(text("UPDATE channel_partner SET credit_limit = 250000 WHERE id = CAST(:p AS uuid)"),
                    {"p": shop.partner})
    await s.commit()
    await s.close()
    admin = await endpoints._as(client, shop, "admin_sales")
    fo = await endpoints._as(client, shop, "field_officer")
    code = (await client.get(f"{V1}/partners/{shop.partner}", headers=admin)).json()["data"]["code"]
    full = {r["Code"]: r for r in _rows(await client.get(f"{V1}/partners/export", headers=admin,
                                                         params={"q": code}))}
    assert full[code]["Credit limit"] == 250000
    seen = {r["Code"]: r for r in _rows(await client.get(f"{V1}/partners/export", headers=fo,
                                                         params={"q": code}))}
    assert code in seen, "a field officer reads the dealer in their district (FS-020)"
    assert seen[code]["Credit limit"] is None, "credit terms need partners.edit"


@pytest.mark.usefixtures("fake_engine")
async def test_subsidy_export_holds_only_the_coordinators_territory(
        client: httpx.AsyncClient, office: subsidy.Office) -> None:
    fo = await subsidy._as(client, office, "field_officer")
    app = await subsidy._app(client, office, fo)
    sc = await subsidy._as(client, office, "state_coordinator")
    rows = _rows(await client.get(f"{V1}/subsidy-applications/export", headers=sc,
                                  params={"q": app["application_no"]}))
    assert [r["Application no"] for r in rows] == [app["application_no"]]
    assert rows[0]["Subsidy"] is not None
    stranger = await subsidy._as(client, office, "stranger")
    assert _rows(await client.get(f"{V1}/subsidy-applications/export", headers=stranger,
                                  params={"q": app["application_no"]})) == []


def test_every_export_takes_every_filter_its_list_takes() -> None:
    """Forwarding is by construction (filters_of reads the list's signature). This
    guards the one way it breaks: a list parameter whose name collides with the
    paging and dependency names the export drops (code review F-3)."""
    from api.main import app
    paths = app.openapi()["paths"]
    paging = {"limit", "cursor", "include_total"}
    for path, ops in paths.items():
        if not path.endswith("/export"):
            continue
        listed = {p["name"] for p in paths[path.removesuffix("/export")]["get"].get("parameters", [])}
        exported = {p["name"] for p in ops["get"].get("parameters", [])}
        # FS-077 EC-13: every export takes `format`, and nothing else beyond the list's own
        assert "format" in exported, path
        assert exported - {"format"} == listed - paging, path


# ── FS-077: CSV and PDF ─────────────────────────────────────────────────────

def _csv(r: httpx.Response) -> list[dict[str, str]]:
    import csv as csv_mod
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"].endswith('.csv"')
    assert r.content.startswith("\ufeff".encode()), "a byte-order mark for Excel"
    return list(csv_mod.DictReader(io.StringIO(r.content.decode("utf-8").lstrip("\ufeff"))))


async def test_a_list_as_csv_holds_the_same_rows_as_excel_and_keeps_scope(
        client: httpx.AsyncClient, shop: Shop, sessions: Sessions) -> None:
    tag = uuid.uuid4().hex[:8]
    fo, dm = shop.ids["field_officer"], shop.ids["district_manager"]
    await _leads(sessions, shop, tag, [fo, fo, dm])
    admin = await endpoints._as(client, shop, "admin_sales")
    q = {"q": f"Export Farmer {tag}"}
    xlsx = _rows(await client.get(f"{V1}/leads/export", headers=admin, params=q))
    rows = _csv(await client.get(f"{V1}/leads/export", headers=admin, params={**q, "format": "csv"}))
    assert sorted(r["Inquiry no"] for r in rows) == sorted(r["Inquiry no"] for r in xlsx)
    assert rows[0]["Estimated value"] == "125000.50", "money as plain two-place text"
    assert rows[0]["Mobile"].startswith("+91"), "a mobile is not formula-guarded (EC-4)"
    h = await endpoints._as(client, shop, "field_officer")
    mine = _csv(await client.get(f"{V1}/leads/export", headers=h, params={**q, "format": "csv"}))
    assert sorted(r["Inquiry no"] for r in mine) == [f"EXP-{tag}-0", f"EXP-{tag}-1"], \
        "the negative: a field officer's CSV never holds the manager's lead"


async def test_pdf_on_a_list_is_refused_before_any_row_is_read(
        client: httpx.AsyncClient, shop: Shop, monkeypatch: pytest.MonkeyPatch) -> None:
    """EC-14: the drain never runs for a refused request."""
    calls: list[int] = []
    real = exports.drain

    async def spy(*a: Any, **k: Any) -> Any:
        calls.append(1)
        return await real(*a, **k)

    monkeypatch.setattr(exports, "drain", spy)
    h = await endpoints._as(client, shop, "admin_sales")
    for path in ("leads", "quotations", "orders", "partners"):
        r = await client.get(f"{V1}/{path}/export", headers=h, params={"format": "pdf"})
        assert r.status_code == 422 and r.json()["error"]["code"] == "format_not_available", (path, r.text)
    assert calls == []
    r = await client.get(f"{V1}/leads/export", headers=h, params={"format": "CSV"})
    assert r.status_code == 422, "formats are lower case"


async def test_every_export_file_is_never_cached(client: httpx.AsyncClient, shop: Shop) -> None:
    """EC-25."""
    h = await endpoints._as(client, shop, "admin_sales")
    for fmt in ("xlsx", "csv"):
        r = await client.get(f"{V1}/partners/export", headers=h, params={"format": fmt, "q": "zz none"})
        assert r.status_code == 200, r.text
        assert r.headers["cache-control"] == "no-store" and r.headers["x-content-type-options"] == "nosniff"


async def test_every_list_export_answers_csv(client: httpx.AsyncClient, shop: Shop) -> None:
    """EC-13: no route silently ignores `format` and answers Excel."""
    h = await endpoints._as(client, shop, "admin_sales")
    for path in ("leads", "quotations", "orders", "complaints", "tasks", "subsidy-applications",
                 "partners", "users"):
        params = {"q": "no such row zz"} if path != "tasks" else {"task_type": "call"}
        _csv(await client.get(f"{V1}/{path}/export", headers=h, params={**params, "format": "csv"}))


@pytest.mark.usefixtures("fake_engine")
async def test_a_subsidy_report_comes_as_csv_and_pdf_and_a_large_pdf_is_refused(
        client: httpx.AsyncClient, office: subsidy.Office, monkeypatch: pytest.MonkeyPatch) -> None:
    from api.routers import exporting
    from api.services import report_pdf
    fo = await subsidy._as(client, office, "field_officer")
    await subsidy._app(client, office, fo)
    sc = await subsidy._as(client, office, "state_coordinator")
    url = f"{V1}/subsidy-reports/stages/export"
    rows = _csv(await client.get(url, headers=sc, params={"format": "csv"}))
    assert rows and "Stage" in rows[0]
    monkeypatch.setattr(report_pdf, "get_settings", lambda: type("S", (), {"pdf_renderer": "html"})())
    r = await client.get(url, headers=sc, params={"format": "pdf"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/html"), "the html renderer labels itself (EC-15)"
    assert r.headers["content-disposition"].endswith('.html"')
    assert "Subsidy stages" in r.text and "Status: open" in r.text
    monkeypatch.setattr(exporting, "PDF_MAX_ROWS", 0)
    r = await client.get(url, headers=sc, params={"format": "pdf"})
    assert r.status_code == 422 and r.json()["error"]["code"] == "export_too_large", r.text
    assert "Use Excel or CSV" in r.json()["error"]["message"]
    assert _csv(await client.get(url, headers=sc, params={"format": "csv"})), "CSV has no PDF cap"
