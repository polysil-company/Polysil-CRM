"""Reports to Excel (GAP-220): the rows, the Total line, the notes and the blank
cells, through the real writer. No database."""

# ruff: noqa: E501  (inline report data)

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from api.services import exports, report_exports


def _ws(report: str, data: dict):
    body = exports.workbook(report, report_exports.columns(report, data), report_exports.rows(report, data))
    ws = load_workbook(io.BytesIO(body)).active
    assert ws is not None
    return ws


def _sheet(report: str, data: dict) -> list[tuple]:
    return [tuple(c.value for c in row) for row in _ws(report, data).iter_rows()]


SALES = {"rows": [{"user": {"id": "u1", "full_name": "Ravi Patel"}, "leads_created": 3, "leads_won": 1,
                   "quotations_sent": 2, "orders": 1, "order_value": "1050.00",
                   "visits": None, "tasks_done": None, "tasks_overdue": None}],
         "truncated": False, "filters": {}, "totals": {"leads_created": 3, "leads_won": 1}}


def test_a_figure_the_caller_may_not_see_is_a_blank_cell_never_zero() -> None:
    """ADR-047: null in JSON stays empty in Excel."""
    rows = _sheet("salesperson-performance", SALES)
    assert rows[0][0] == "Person" and rows[0][6] == "Visits"
    ravi = rows[1]
    assert ravi[0] == "Ravi Patel"
    assert isinstance(ravi[5], (int, float)) and Decimal(str(ravi[5])) == Decimal("1050"), "money is a number"
    assert ravi[6] is None and ravi[7] is None and ravi[8] is None


def test_the_total_line_carries_only_what_the_report_totals() -> None:
    total = _sheet("salesperson-performance", SALES)[2]
    assert total[0] == "Total" and total[1] == 3 and total[2] == 1
    assert all(v is None for v in total[3:]), "no invented totals for untotalled columns"


def test_a_truncated_report_says_so_below_the_rows() -> None:
    data = {**SALES, "truncated": True}
    assert _sheet("salesperson-performance", data)[-1][0].startswith("Only the first 1,000 rows")


def test_grouped_reports_name_their_group_and_flatten_lost_stages() -> None:
    lost = {"group_by": "reason", "truncated": False, "filters": {}, "totals": {"lost": 2},
            "rows": [{"key": {"id": "r", "label": "Price"}, "count": 2, "share_pct": "100.0",
                      "avg_days_to_loss": "4.5", "by_stage": {"qualified": 2}}]}
    rows = _sheet("lost-leads", lost)
    assert rows[0][0] == "Lost reason" and "Lost from qualified" in rows[0]
    assert rows[1][rows[0].index("Lost from qualified")] == 2
    assert rows[1][rows[0].index("Lost from new")] is None
    assert rows[2][:2] == ("Total", 2)


def test_complaint_refunds_are_a_note_and_a_formula_stays_text() -> None:
    data = {"group_by": "type", "truncated": False, "filters": {},
            "totals": {"count": 1, "resolved": 1, "refunds": {"count": 1, "amount": "12500.00"}},
            "rows": [{"key": {"id": "t", "label": "=HYPERLINK(\"x\")"}, "count": 1, "resolved": 1,
                      "resolved_within_sla_pct": "100.0", "response_breaches": 0}]}
    ws = _ws("complaints", data)
    rows = [tuple(c.value for c in row) for row in ws.iter_rows()]
    assert rows[1][0] == '=HYPERLINK("x")'
    # code review: the value reads the same for a live formula; the type does not
    assert ws["A2"].data_type == "s", "stored as text, never a live formula"
    assert rows[3][0] == "Refunds paid: 1, amount 12500.00"


@pytest.mark.parametrize("report", report_exports.NAMES)
def test_every_report_has_columns(report: str) -> None:
    cols = report_exports.columns(report, {"group_by": "owner"})
    assert cols and len({c.header for c in cols}) == len(cols)
