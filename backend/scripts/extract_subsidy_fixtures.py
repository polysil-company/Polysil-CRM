"""Extract the golden fixtures for the subsidy engine from the client's workbooks.

    python scripts/extract_subsidy_fixtures.py            # writes tests/fixtures/subsidy/*.json
    python scripts/extract_subsidy_fixtures.py --force    # overwrite a fixture whose inputs changed

Import, never transcribe (FS-008 section 5). Every input the calculation needs is
read from the sheet the designer fills in, and every expected output is the
workbook's cached value for a named cell, tagged `rounded` when the cell's own
formula rounds it and `raw` when it does not. The fixture files are committed:
`data/` is not in the shared repository, so the suite must run without it.

The three workbooks are named here, not globbed: `data/boq/` also holds a `(1)`
copy of the Mini Sprinkler file and an `.xls` of the Sprinkler file.
"""

# ruff: noqa: E501  (cell maps read better on one line, as the migrations' pasted SQL does)

from __future__ import annotations

import argparse
import json
import re
import sys
import warnings
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")

import openpyxl  # noqa: E402  (a dev dependency; the warning filter must precede the import)

ROOT = Path(__file__).resolve().parents[1]
BOQ = ROOT / "data" / "boq"
OUT = ROOT / "tests" / "fixtures" / "subsidy"

FILES = {
    "drip": "POLYSIL Drip New BOQ Two Crop GST 5% 13-06-2026.xlsx",
    "mini_sprinkler": "POLYSIL Mini Sprinkler BOQ 13-06-2026 5 % GST.xlsx",
    "sprinkler": "POLYSIL SPRINKLER  BOQ -13-06-2026 GST 5 %  BRASS & PLASTIC.xlsx",
}

_ROUND = re.compile(r"ROUND\(", re.I)


class Book:
    def __init__(self, path: Path) -> None:
        self.formulas = openpyxl.load_workbook(path, data_only=False)
        self.values = openpyxl.load_workbook(path, data_only=True)

    def value(self, sheet: str, cell: str) -> Any:
        return self.values[sheet][cell].value

    def formula(self, sheet: str, cell: str) -> str | None:
        v = self.formulas[sheet][cell].value
        return v if isinstance(v, str) and v.startswith("=") else None

    def expected(self, sheet: str, cell: str) -> dict[str, Any]:
        """A named output: its cached value, and whether its own formula rounds."""
        f = self.formula(sheet, cell)
        v = self.value(sheet, cell)
        return {"cell": f"'{sheet}'!{cell}", "value": _num(v), "kind": "rounded" if f and _ROUND.search(f) else "raw",
                "formula": f}

    def rows_between(self, sheet: str, col_marker: str, start_text: str, end_text: str) -> list[int]:
        """Row numbers strictly between the row whose `col_marker` cell starts with
        `start_text` and the next whose cell starts with `end_text`."""
        ws = self.values[sheet]
        start = end = None
        for r in range(1, ws.max_row + 1):
            t = ws[f"{col_marker}{r}"].value
            if not isinstance(t, str):
                continue
            if start is None and t.strip().upper().startswith(start_text.upper()):
                start = r
            elif start is not None and t.strip().upper().startswith(end_text.upper()):
                end = r
                break
        assert start is not None and end is not None, (sheet, start_text, end_text)
        return list(range(start + 1, end))


def _num(v: Any) -> Any:
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, (int, float)):
        return repr(float(v)) if isinstance(v, float) else v
    return str(v)


SENTINEL = "Select Crop here"


def _crop(v: Any) -> str | None:
    """The workbook's empty-crop sentinel is null on the wire."""
    if v is None or str(v).strip() == "" or str(v).strip() == SENTINEL:
        return None
    return str(v).strip()


def _dec(v: Any) -> str:
    """An input as the API takes it: a decimal string."""
    if v is None or v == "":
        return "0"
    return repr(float(v)) if isinstance(v, float) else str(v)


def _lines(book: Book, sheet: str, rows: list[int], qty_col: str, rate_col: str = "H",
           desc_col: str = "C", uom_col: str = "G") -> list[dict[str, str]]:
    out = []
    for r in rows:
        desc = book.value(sheet, f"{desc_col}{r}")
        qty = book.value(sheet, f"{qty_col}{r}")
        rate = book.value(sheet, f"{rate_col}{r}")
        if not isinstance(desc, str) or not desc.strip() or qty in (None, 0, "") or rate in (None, ""):
            continue
        out.append({"row": r, "description": desc.strip(), "uom": str(book.value(sheet, f"{uom_col}{r}") or ""),
                    "rate": _dec(rate), "qty": _dec(qty)})
    return out


def drip(book: Book) -> dict[str, Any]:
    qa, qs, ns1, ns2, j1, j2, qc = ("Q A  for one Farmer", "Quo Summary", "New Subsidy Calculation C1",
                                    "New Subsidy Calculation C2", "7 Year Jantri Calculation C1",
                                    "7 Year Jantri Calculation C2", "Q C & F for one farmer ")
    head_rows = book.rows_between(qa, "C", "HEAD UNIT", "Sub Total A")
    field_rows = book.rows_between(qa, "C", "FIELD UNIT", "Sub Total (B)")
    inputs = {
        "system_type": "drip",
        "group_total_area": _dec(book.value(qa, "M4")),
        "sump": {"rate_per_ha": _dec(book.value(qa, "G71")), "label": str(book.value(qa, "C71"))},
        "installation_rate_per_ha": _dec(book.value(qa, "H67")),
        "education_amount": _dec(book.value(qc, "F16")),
        "head_lines": _lines(book, qa, head_rows, qty_col="M"),
        "crops": [
            {"crop": _crop(book.value(qa, "I6")), "inter_crop": _crop(book.value(qa, "J6")),
             "area": _dec(book.value(qa, "I7")), "crop_spacing": str(book.value(qa, "I8")),
             "lateral_spacing": _dec(book.value(qa, "I9")), "lines": _lines(book, qa, field_rows, qty_col="I")},
            {"crop": _crop(book.value(qa, "K6")), "inter_crop": _crop(book.value(qa, "L6")),
             "area": _dec(book.value(qa, "K7")), "crop_spacing": str(book.value(qa, "K8")),
             "lateral_spacing": _dec(book.value(qa, "K9")), "lines": _lines(book, qa, field_rows, qty_col="K")},
        ],
    }
    expected: dict[str, Any] = {}
    for col, key in (("I", "crop1"), ("J", "crop2"), ("K", "total")):
        for r in range(18, 41):
            label = book.value(qs, f"D{r}") or book.value(qs, f"B{r}")
            if book.value(qs, f"{col}{r}") is None:
                continue
            expected[f"summary.{key}.{col}{r}"] = {**book.expected(qs, f"{col}{r}"), "label": str(label)}
    for r in range(29, 37):
        expected[f"categories.crop1.{r}"] = {"name": str(book.value(qs, f"N{r}")),
                                             "subsidy": book.expected(qs, f"Q{r}"),
                                             "farmer_share": book.expected(qs, f"R{r}")}
    for r in range(50, 58):
        expected[f"categories.crop2.{r}"] = {"name": str(book.value(qs, f"N{r}")),
                                             "subsidy": book.expected(qs, f"Q{r}"),
                                             "farmer_share": book.expected(qs, f"R{r}")}
    for sheet, key in ((ns1, "crop1"), (ns2, "crop2")):
        for cell, name in (("C30", "regular"), ("C31", "seven_year"), ("C32", "regular_with_sump"),
                           ("C55", "standard_spacing"), ("C56", "spacing_for_subsidy"), ("C29", "cost")):
            expected[f"jantri.{key}.{name}"] = book.expected(sheet, cell)
        for r in range(34, 40):
            expected[f"capped.{key}.{r}"] = {"label": str(book.value(sheet, f"A{r}")),
                                             "formula": book.formula(sheet, f"C{r}"),
                                             "value": book.expected(sheet, f"C{r}")}
    for sheet, key in ((j1, "crop1"), (j2, "crop2")):
        expected[f"seven_year.{key}.unit_cost"] = book.expected(sheet, "C34")
        expected[f"seven_year.{key}.spacing"] = book.expected(sheet, "C33")
        expected[f"seven_year.{key}.area"] = book.expected(sheet, "C32")
    return {"inputs": inputs, "expected": expected}


def mini(book: Book) -> dict[str, Any]:
    qa, qs, ns, j7 = "Quo A", "Quo Summary", "New Subsidy Calculation", "7 Year Jantri Calculation"
    head_rows = book.rows_between(qa, "C", "HEAD UNIT", "Sub Total")
    field_rows = book.rows_between(qa, "C", "FIELD UNIT", "Sub Total (B)")
    inputs = {
        "system_type": "mini_sprinkler",
        "group_total_area": _dec(book.value(qa, "I5")),
        "sump": {"rate_per_ha": _dec(book.value(qa, "G70")), "label": str(book.value(qa, "C70"))},
        "installation_rate_per_ha": _dec(book.value(qa, "H66")),
        "education_amount": _dec(book.value(qa, "J65")),
        "head_lines": _lines(book, qa, head_rows, qty_col="I"),
        "crops": [
            {"crop": _crop(book.value(qa, "I8")), "inter_crop": None,
             "area": _dec(book.value(qa, "I9")), "crop_spacing": str(book.value(qa, "I10")),
             "lateral_spacing": _dec(book.value(qa, "I11")), "lines": _lines(book, qa, field_rows, qty_col="I")},
        ],
    }
    expected: dict[str, Any] = {}
    for r in range(18, 41):
        if book.value(qs, f"I{r}") is None:
            continue
        expected[f"summary.crop1.I{r}"] = {**book.expected(qs, f"I{r}"),
                                           "label": str(book.value(qs, f"D{r}") or book.value(qs, f"B{r}"))}
    expected["jantri.crop1.regular_prorated"] = book.expected(qs, "M31")
    for r in range(33, 41):
        expected[f"categories.crop1.{r}"] = {"name": str(book.value(qs, f"L{r}")),
                                             "subsidy": book.expected(qs, f"O{r}"),
                                             "farmer_share": book.expected(qs, f"P{r}"),
                                             "subsidy_pct": book.expected(qs, f"Q{r}")}
    for r in range(43, 49):
        expected[f"gsdma.crop1.{r}"] = {"name": str(book.value(qs, f"L{r}")),
                                        "gsdma_farmer_share": book.expected(qs, f"Q{r}")}
    for cell, name in (("G30", "regular"), ("G31", "seven_year"), ("G32", "regular_with_sump"), ("G29", "cost")):
        expected[f"jantri.crop1.{name}"] = book.expected(ns, cell)
    expected["seven_year.crop1.unit_cost"] = book.expected(j7, "G34")
    expected["seven_year.crop1.spacing"] = book.expected(j7, "G33")
    return {"inputs": inputs, "expected": expected}


def sprinkler(book: Book) -> dict[str, Any]:
    boq, fsc = "BOQ", "Farmer Share Calculation"
    inputs = {
        "system_type": "sprinkler",
        "crops": [{"crop": _crop(book.value(boq, "J10")), "inter_crop": None, "area": _dec(book.value(boq, "J11")),
                   "crop_spacing": str(book.value(boq, "J12")), "lateral_spacing": _dec(book.value(boq, "J13")),
                   "lines": []}],
        "pipe_size_mm": int(book.value(boq, "A2")),
        "nozzle": "plastic" if "Plastic" in str(book.value(boq, "C18")) else "brass",
    }
    expected: dict[str, Any] = {}
    for r in range(16, 26):
        if book.value(boq, f"C{r}") is None:
            continue
        expected[f"lines.{r}"] = {"description": str(book.value(boq, f"C{r}")), "uom": str(book.value(boq, f"G{r}")),
                                  "qty": book.expected(boq, f"H{r}"), "rate": book.expected(boq, f"I{r}"),
                                  "amount": book.expected(boq, f"J{r}")}
    for cell in ("J24", "J26", "J27", "J28", "J29"):
        expected[f"boq.{cell}"] = {**book.expected(boq, cell), "label": str(book.value(boq, f"B{cell[1:]}"))}
    for r in range(20, 32):
        if book.value(boq, f"R{r}") is None:
            continue
        expected[f"summary.R{r}"] = {**book.expected(boq, f"R{r}"), "label": str(book.value(boq, f"M{r}"))}
    expected["inspection.computed"] = book.expected(boq, "Q46")
    expected["inspection.floor"] = book.expected(boq, "Q47")
    expected["jantri.regular"] = book.expected(fsc, "C4")
    expected["jantri.seven_year"] = book.expected(fsc, "C17")
    for r in range(6, 12):
        expected[f"categories.regular.{r}"] = {"name": str(book.value(fsc, f"A{r}")),
                                              "subsidy": book.expected(fsc, f"B{r}"),
                                              "farmer_share": book.expected(fsc, f"C{r}"),
                                              "subsidy_pct": book.expected(fsc, f"D{r}"),
                                              "dbt_farmer_payable": book.expected(fsc, f"E{r}")}
    for r in (19, 20):
        expected[f"categories.seven_year.{r}"] = {"name": str(book.value(fsc, f"A{r}")),
                                                 "subsidy": book.expected(fsc, f"B{r}"),
                                                 "farmer_share": book.expected(fsc, f"C{r}"),
                                                 "subsidy_pct": book.expected(fsc, f"D{r}")}
    return {"inputs": inputs, "expected": expected}


EXTRACTORS = {"drip": drip, "mini_sprinkler": mini, "sprinkler": sprinkler}


def _grid(book: Book, sheet: str, row_label_col: str, first_row: int, last_row: int,
          header_row: int, first_col: str, last_col: str) -> dict[str, Any]:
    """A 2-D unit-cost matrix: spacing rows by area columns, cells at the workbook's precision."""
    ws = book.values[sheet]
    from openpyxl.utils import column_index_from_string, get_column_letter
    c0, c1 = column_index_from_string(first_col), column_index_from_string(last_col)
    areas = [_num(ws[f"{get_column_letter(c)}{header_row}"].value) for c in range(c0, c1 + 1)]
    rows = []
    for r in range(first_row, last_row + 1):
        spacing = ws[f"{row_label_col}{r}"].value
        if spacing is None:
            continue
        rows.append({"spacing": _num(spacing),
                     "costs": [_num(ws[f"{get_column_letter(c)}{r}"].value) for c in range(c0, c1 + 1)]})
    return {"source": f"'{sheet}'!{first_col}{first_row}:{last_col}{last_row}", "areas": areas, "rows": rows}


def _table_1d(book: Book, sheet: str, first_row: int, last_row: int, area_col: str, cost_col: str) -> dict[str, Any]:
    ws = book.values[sheet]
    return {"source": f"'{sheet}'!{area_col}{first_row}:{cost_col}{last_row}",
            "rows": [{"area": _num(ws[f"{area_col}{r}"].value), "cost": _num(ws[f"{cost_col}{r}"].value)}
                     for r in range(first_row, last_row + 1) if ws[f"{area_col}{r}"].value is not None]}


# The Sprinkler quantity matrix by row: two nozzle rows (brass, plastic) share a
# component name, so the code is by row, as the BOQ's lines are.
_COMPONENT_CODES = {7: "pipe", 8: "coupler", 9: "nozzle_brass", 10: "nozzle_plastic", 11: "riser",
                    12: "pcn", 13: "bend", 14: "tee", 15: "end_plug", 16: "transport"}


def matrices(books: dict[str, Book]) -> dict[str, Any]:
    """Every master the engine reads, from its named range (FS-008 section 5)."""
    drip_b, mini_b, spr_b = books["drip"], books["mini_sprinkler"], books["sprinkler"]
    ws = spr_b.values["SPRINKLER 13.06.2026"]
    from openpyxl.utils import get_column_letter
    areas = [_num(ws[f"{get_column_letter(c)}4"].value) for c in range(5, 30)]   # E4:AC4
    components = []
    for r in range(7, 17):
        name = ws[f"C{r}"].value
        if not name:
            continue
        components.append({"row": r, "code": _COMPONENT_CODES[r], "component": str(name).strip(),
                           "description": str(ws[f"B{r}"].value or "").strip(), "uom": str(ws[f"D{r}"].value),
                           "qty": [_num(ws[f"{get_column_letter(c)}{r}"].value) for c in range(5, 30)]})
    crops = drip_b.values["New Subsidy Calculation C1"]
    crop_rows = [{"crop": str(crops[f"B{r}"].value).strip(), "standard_spacing": _num(crops[f"C{r}"].value)}
                 for r in range(59, 138) if crops[f"B{r}"].value not in (None, "")]
    return {
        "drip": {"regular": _grid(drip_b, "New Subsidy Calculation C1", "B", 4, 17, 3, "C", "I"),
                 "seven_year": _grid(drip_b, "7 Year Jantri Calculation C1", "B", 4, 16, 3, "C", "I")},
        "mini_sprinkler": {"regular": _grid(mini_b, "New Subsidy Calculation", "L", 4, 6, 3, "M", "S"),
                           "seven_year": _grid(mini_b, "7 Year Jantri Calculation", "L", 4, 5, 3, "M", "S")},
        "sprinkler": {"regular": _table_1d(spr_b, "Farmer Share Calculation", 2, 26, "L", "M"),
                      "seven_year": _table_1d(spr_b, "Farmer Share Calculation", 30, 54, "L", "M"),
                      "quantities": {"source": "'SPRINKLER 13.06.2026'!B7:AC16", "areas": areas, "components": components},
                      "pipe_size_bands": {"75": "areas up to 2.0", "90": "areas from 2.01", "source": "'SPRINKLER 13.06.2026'!E6, N6"}},
        "crops": {"source": "'New Subsidy Calculation C1'!B59:C137", "rows": crop_rows},
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--force", action="store_true", help="overwrite a fixture whose inputs changed")
    args = parser.parse_args(argv)
    OUT.mkdir(parents=True, exist_ok=True)
    books: dict[str, Book] = {}
    for key, name in FILES.items():
        path = BOQ / name
        if not path.exists():
            print(f"{key}: {path.name} not found under data/boq; skipped")
            continue
        books[key] = Book(path)
        fixture = EXTRACTORS[key](books[key])
        fixture["source"] = name
        target = OUT / f"{key}.json"
        if target.exists() and not args.force:
            old = json.loads(target.read_text(encoding="utf-8"))
            if old.get("inputs") != fixture["inputs"]:
                print(f"{key}: the inputs changed; rerun with --force to overwrite")
                return 1
        target.write_text(json.dumps(fixture, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"{key}: {len(fixture['expected'])} expected cells -> {target.relative_to(ROOT)}")
    if len(books) == len(FILES):
        target = OUT / "matrices.json"
        target.write_text(json.dumps(matrices(books), indent=2, ensure_ascii=False) + "\n",
                          encoding="utf-8")
        print(f"matrices -> {target.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
