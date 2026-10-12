"""Extract other states' subsidy masters and worked examples from the client's
quotation formats (FS-056).

    python scripts/extract_state_masters.py

Reads `data/quotation-formats/{ap,cg,rj,up}/*.xlsx` in the main checkout (the
formats are not in the shared repository) and writes two files:

* `tests/fixtures/subsidy/states.json`: public scheme data per "<STATE>/<system>":
  the engine settings, parameters, categories and unit cost tables.
* `tests/fixtures/subsidy/states_samples.json`: company pricing and the cached
  worked example per workbook: component rates, quantity tables, inputs and the
  expected figures.

Numbers are decimal strings. A cell that is a constant product (`=4250*1.27`) is
recomputed in Decimal from its constants. Any other number is the cell's value;
a float is read through `repr` and then cut to 15 significant digits, the
precision Excel itself keeps, so `108254.09999999999` is `108254.1`. Area keys
are rounded to 3 dp ("0.4", "1", "2.01").
"""

# ruff: noqa: E501

from __future__ import annotations

import json
import re
import sys
import warnings
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")

import openpyxl  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA_CANDIDATES = (ROOT / "data" / "quotation-formats", Path(r"D:\Polysil-CRM\data\quotation-formats"))
OUT = ROOT / "tests" / "fixtures" / "subsidy"

FILES = {
    "AP/drip": "ap/AP-Drip-Portal-BOQ.xlsx",
    "AP/sprinkler": "ap/AP-Sprinkler-Quotation-Format.xlsx",
    "CG/drip": "cg/CG-Drip-Quotation-Format.xlsx",
    "CG/sprinkler": "cg/CG-Sprinkler-Quotation-Format.xlsx",
    "RJ/drip": "rj/RJ-Drip-Quotation-Format.xlsx",
    "RJ/mini_sprinkler": "rj/RJ-Mini-Sprinkler-Quotation-Format.xlsx",
    "RJ/sprinkler": "rj/RJ-Sprinkler-Quotation-Format.xlsx",
    "UP/drip": "up/UP-Drip-Quotation-Format.xlsx",
    "UP/mini_sprinkler": "up/UP-Mini-Sprinkler-Quotation-Format.xlsx",
    "UP/sprinkler": "up/UP-Sprinkler-Quotation-Format.xlsx",
}

_PRODUCT = re.compile(r"^=\s*([0-9.]+)\s*([*/])\s*([0-9.]+)\s*$")


# ── numbers ──────────────────────────────────────────────────────────────────

def dec(v: Any) -> Decimal:
    """A cell value as a Decimal: ints exactly, floats via repr cut to Excel's 15
    significant digits."""
    if isinstance(v, bool) or v is None:
        raise ValueError(f"not a number: {v!r}")
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        d = Decimal(repr(v))
        if d == 0:
            return Decimal(0)
        with localcontext() as ctx:
            ctx.prec = 15
            return +d
    return Decimal(str(v).strip())


def s(d: Decimal | int | float) -> str:
    """A decimal string without exponent or trailing zeros."""
    d = d if isinstance(d, Decimal) else dec(d)
    if d == d.to_integral_value():
        return str(d.quantize(Decimal(1)))
    out = format(d.normalize(), "f")
    return out


def area_key(v: Any) -> str:
    return s(dec(v).quantize(Decimal("0.001")))


class Book:
    def __init__(self, path: Path) -> None:
        self.f = openpyxl.load_workbook(path, data_only=False)
        self.v = openpyxl.load_workbook(path, data_only=True)

    def raw(self, sheet: str, cell: str) -> Any:
        return self.v[sheet][cell].value

    def text(self, sheet: str, cell: str) -> str | None:
        v = self.raw(sheet, cell)
        return None if v is None else str(v)

    def num(self, sheet: str, cell: str, default: Decimal | None = None) -> Decimal:
        """Recompute a constant product from its constants; else the cached value."""
        f = self.f[sheet][cell].value
        if isinstance(f, str):
            m = _PRODUCT.match(f)
            if m:
                a, op, b = Decimal(m.group(1)), m.group(2), Decimal(m.group(3))
                return a * b if op == "*" else a / b
        v = self.raw(sheet, cell)
        if v is None or v == "":
            if default is not None:
                return default
            raise ValueError(f"{sheet}!{cell} is empty")
        return dec(v)


def cat(code: str, name: str, pct: int, order: int) -> dict[str, Any]:
    return {"code": code, "name": name, "pct": str(pct), "variant": "regular", "sort_order": order}


def system(pipeline: str, jantri: str, qty: str, cap: str, blocks: list[str], lines: list[str], *,
           outside: str = "clamp", head: bool = False, group: bool = False) -> dict[str, Any]:
    return {"pipeline": pipeline, "jantri_variant": jantri, "quantity_source": qty, "cap_mode": cap,
            "rounded_blocks": blocks, "rounded_lines": lines, "spacing_rule": "design",
            "spacing_outside_table": outside, "has_head_unit": head, "supports_group": group,
            "crop_count_max": 1}


def example_input(**kw: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "crops": [], "head_lines": [], "group_total_area": None, "installation_rate_per_ha": "0",
        "sump_rate_per_ha": "0", "nozzle": None, "pipe_size_mm": None, "spacing_label": None,
        "hydrocyclone_m3": "0", "gravel_filter_m3": "0", "head_unit_existing": False,
        "pvc_existing": False, "guided_cost": None, "fittings_amount": None,
    }
    base.update(kw)
    return base


def crop(area: Decimal, lateral: Decimal | None, lines: list[dict[str, str]]) -> dict[str, Any]:
    return {"crop": None, "area": s(area), "lateral_spacing": None if lateral is None else s(lateral),
            "lines": lines}


def lines_of(b: Book, sheet: str, rows: range, *, desc: str, uom: str, rate: str, qty: str) -> list[dict[str, str]]:
    """Every quoted line with a description and a rate; an empty quantity is 0."""
    out = []
    for r in rows:
        d = b.text(sheet, f"{desc}{r}")
        rv = b.raw(sheet, f"{rate}{r}")
        if not d or rv is None or isinstance(rv, str):
            continue
        out.append({"description": d.strip(), "uom": (b.text(sheet, f"{uom}{r}") or "").strip() or "No.",
                    "rate": s(b.num(sheet, f"{rate}{r}")),
                    "qty": s(b.num(sheet, f"{qty}{r}", Decimal(0)))})
    return out


def expected(total: Decimal, unit_cost: Decimal, subsidy: Decimal, share: Decimal) -> dict[str, str]:
    return {"total_incl_gst": s(total), "unit_cost": s(unit_cost), "subsidy": s(subsidy),
            "farmer_share": s(share)}


# ── drip and mini sprinkler: UP (seven blocks, bilinear) ─────────────────────

UP_CATS = [cat("big", "Big farmer (80%)", 80, 1), cat("small", "Small farmer (90%)", 90, 2)]
UP_PARAMS = {"insurance_rate": "0.00169", "education_amount": "0", "inspection_floor": "200",
             "inspection_rate": "0.004", "gst_material_half": "0.025", "gst_service_half": "0.09",
             "min_area_prorate": "0", "max_area_scaling": "0", "jantri_guards": "0",
             "seven_year_area_min": "0.2", "seven_year_area_max": "5"}
UP_BLOCKS = ["total_abc_gst", "insurance", "cost_of_mis", "total_gst"]


def up_bilinear(b: Book, rows: range) -> dict[str, Any]:
    sh = "Unit Cost and Calculation"
    areas = [area_key(b.raw(sh, f"{c}3")) for c in "CDEFGHI"]
    out = [{"spacing": s(b.num(sh, f"B{r}")), "costs": [s(b.num(sh, f"{c}{r}")) for c in "CDEFGHI"]}
           for r in rows]
    out.sort(key=lambda x: Decimal(x["spacing"]), reverse=True)
    return {"kind": "bilinear_2d", "areas": areas, "rows": out}


def up_drip(b: Book) -> tuple[dict, dict]:
    q, calc = "Drip Quotation-Up", "Unit Cost and Calculation"
    pub = {"scheme": "UP", "system_type": "drip",
           "system": system("full_abcdefg", "bilinear_2d", "manual", "incl_gst", UP_BLOCKS, [],
                            outside="extrapolate", head=True, group=True),
           "parameters": dict(UP_PARAMS), "categories": UP_CATS, "unit_cost": up_bilinear(b, range(4, 17))}
    inp = example_input(
        crops=[crop(b.num(q, "E4"), b.num(q, "E7"),
                    lines_of(b, q, range(24, 40), desc="B", uom="C", rate="D", qty="E"))],
        head_lines=lines_of(b, q, range(11, 22), desc="B", uom="C", rate="D", qty="E"),
        group_total_area=s(b.num(q, "F2")), installation_rate_per_ha=s(b.num(q, "D43")),
        sump_rate_per_ha=s(b.num(q, "D44")))
    # 'Big Farmer, 80 %' is C39:D39
    sample = {"component_rates": [], "quantities": None,
              "example": {"input": inp, "category": "big",
                          "expected": expected(b.num(q, "F69"), b.num(calc, "C37"),
                                               b.num(calc, "C39"), b.num(calc, "D39"))}}
    assert b.text(q, "E8") == "Big Farmer, 80 %"
    return pub, sample


def up_mini(b: Book) -> tuple[dict, dict]:
    q, calc = "Mini Sprinkler", "Unit Cost and Calculation"
    pub = {"scheme": "UP", "system_type": "mini_sprinkler",
           "system": system("full_abcdefg", "bilinear_2d", "manual", "incl_gst", UP_BLOCKS, [],
                            outside="extrapolate", head=True, group=True),
           "parameters": dict(UP_PARAMS), "categories": UP_CATS, "unit_cost": up_bilinear(b, range(4, 8))}
    inp = example_input(
        crops=[crop(b.num(q, "E4"), b.num(q, "E6"),
                    lines_of(b, q, range(23, 51), desc="B", uom="C", rate="D", qty="E"))],
        head_lines=lines_of(b, q, range(10, 21), desc="B", uom="C", rate="D", qty="E"),
        group_total_area=s(b.num(q, "F2")), installation_rate_per_ha=s(b.num(q, "D54")),
        sump_rate_per_ha=s(b.num(q, "D55")))
    assert b.text(q, "E7") == "Small Farmer, 90 %"
    sample = {"component_rates": [], "quantities": None,
              "example": {"input": inp, "category": "small",
                          "expected": expected(b.num(q, "F81"), b.num(calc, "C26"),
                                               b.num(calc, "C29"), b.num(calc, "D29"))}}
    return pub, sample


# ── sprinkler: CG, RJ, UP (set size x area) ──────────────────────────────────

QTY_CODES = {2: "pipe", 3: "coupler", 4: "nozzle", 5: "riser", 6: "pcn", 7: "bend", 8: "tee",
             9: "end_plug", 10: "transport"}
PRICE_CODES = (("hdpe pipes", "pipe"), ("sprinkler coupler", "coupler"), ("riser", "riser"),
               ("pump conecting", "pcn"), ("bend", "bend"), ("tee", "tee"), ("end plug", "end_plug"))


def size_mm(label: Any) -> int:
    return int(str(label).upper().replace("MM", "").strip())


def sprinkler_rates(b: Book) -> list[dict[str, Any]]:
    sh = "Price List"
    rates: list[dict[str, Any]] = []
    sizes: list[int] = []
    riser = None
    for r in range(2, 30):
        d, rate = b.text(sh, f"B{r}"), b.raw(sh, f"C{r}")
        if not d:
            continue
        low = d.lower()
        if low.startswith("secondary transportation"):
            rates.append({"code": "transport", "description": d, "uom": "Ha",
                          "rate": s(b.num(sh, f"C{r}", Decimal(0))), "pipe_size_mm": None, "nozzle": None})
            continue
        if rate is None:
            continue                              # 110 mm has no rate in CG and RJ
        if low.startswith("sprinkler nozzles"):
            noz = "brass" if "brass" in low else "plastic"
            rates.append({"code": "nozzle", "description": d, "uom": "No.", "rate": s(b.num(sh, f"C{r}")),
                          "pipe_size_mm": None, "nozzle": noz})
            continue
        code = next(c for p, c in PRICE_CODES if low.startswith(p))
        if code == "riser":
            riser = (d, b.num(sh, f"C{r}"))
            continue
        size = size_mm(b.raw(sh, f"A{r}"))
        if code == "pipe":
            sizes.append(size)
        rates.append({"code": code, "description": d, "uom": "No.", "rate": s(b.num(sh, f"C{r}")),
                      "pipe_size_mm": size, "nozzle": None})
    # The riser is 20 mm whatever the set; the engine reads it by the set size,
    # as the GGRC masters hold it (one row per size, migration 009).
    assert riser is not None
    for size in sizes:
        rates.append({"code": "riser", "description": riser[0], "uom": "No.", "rate": s(riser[1]),
                      "pipe_size_mm": size, "nozzle": None})
    return rates


def sprinkler_quantities(b: Book, nozzle: str) -> dict[str, Any]:
    sh = "Area wise QTY."
    cols = [c for c in range(4, 29) if b.raw(sh, f"{openpyxl.utils.get_column_letter(c)}1") is not None]
    letters = [openpyxl.utils.get_column_letter(c) for c in cols]
    areas = [area_key(b.raw(sh, f"{L}1")) for L in letters]
    rows: dict[str, dict[str, str]] = {}
    for r, code in QTY_CODES.items():
        key = f"nozzle_{nozzle}" if code == "nozzle" else code
        rows[key] = {a: s(b.num(sh, f"{L}{r}")) for a, L in zip(areas, letters, strict=True)}
    return {"areas": areas, "rows": rows}


def sprinkler_unit_cost(b: Book) -> dict[str, Any]:
    sh = "Unit Cost & Calculation"
    rows: dict[str, dict[str, str]] = {}
    for c in "BCDE":
        size = str(size_mm(b.raw(sh, f"{c}3")))
        cells = {}
        for r in range(4, 29):
            v = b.raw(sh, f"{c}{r}")
            if v is not None and not isinstance(v, str):
                cells[area_key(b.raw(sh, f"A{r}"))] = s(b.num(sh, f"{c}{r}"))
        if cells:
            rows[size] = cells
    return {"kind": "rows", "rows": rows}


def sprinkler(b: Book, state: str) -> tuple[dict, dict]:
    calc = "Unit Cost & Calculation"
    q = {"CG": "Quotation SIS-CG", "AP": "Quotation SIS-APMIP"}.get(state, "Quotation SIS-UP")
    nozzle = "brass" if "brass" in (b.text(q, "B16" if state == "AP" else "B14") or "").lower() else "plastic"
    if state == "AP":
        return ap_sprinkler(b, q, calc, nozzle)
    if state == "CG":
        sys_ = system("gst_inclusive", "size_area_2d", "area_matrix", "incl_gst", [], ["field", "transport"])
        params = {"pipe_size_from_input": "1", "share_floor_zero": "1",
                  "flat_inspection": s(b.num(q, "J22")), "gst_inclusive_rate": "0.05",
                  "exact_area_match_required": "0"}
        cats = [cat("small", "Small farmer (55%)", 55, 1), cat("big", "Big farmer (45%)", 45, 2)]
        total = b.num(q, "J21")                      # the cap's figure, before the flat inspection
    else:
        sys_ = system("field_inspection_only", "size_area_2d", "area_matrix", "incl_gst", [],
                      ["field", "transport"])
        params = {"pipe_size_from_input": "1", "share_floor_zero": "1" if state == "RJ" else "0",
                  "inspection_floor": "200", "inspection_rate": "0.004", "gst_material_half": "0.025",
                  "gst_service_half": "0.09", "exact_area_match_required": "0"}
        if state == "RJ":
            cats = [cat("small_marginal", "Small or marginal farmer (75%)", 75, 1),
                    cat("sc", "SC farmer (75%)", 75, 2), cat("st", "ST farmer (75%)", 75, 3),
                    cat("woman", "Woman farmer (75%)", 75, 4), cat("others", "Others farmer (70%)", 70, 5)]
        else:
            cats = [cat("small", "Small farmer (75%)", 75, 1), cat("other", "Other farmer (65%)", 65, 2)]
        total = b.num(q, "R28")
    pub = {"scheme": state, "system_type": "sprinkler", "system": sys_, "parameters": params,
           "categories": cats, "unit_cost": sprinkler_unit_cost(b)}
    chosen = b.text(q, "H8")
    labels = {r: b.text(calc, f"H{r}") for r in range(8, 13)}
    row = next(r for r, t in labels.items() if t == chosen)
    code = cats[row - 8]["code"]
    inp = example_input(crops=[crop(b.num(q, "H6"), None, [])], nozzle=nozzle,
                        pipe_size_mm=size_mm(b.raw(q, "H5")))
    sample = {"component_rates": sprinkler_rates(b), "quantities": sprinkler_quantities(b, nozzle),
              "example": {"input": inp, "category": code,
                          "expected": expected(total, b.num(calc, "I6"), b.num(calc, f"J{row}"),
                                               b.num(calc, f"K{row}"))}}
    return pub, sample


# ── AP: sprinkler (excl GST, ceiling) and drip (the portal's figures) ───────

_CEILING = re.compile(r"MIN\(MIN\(I6,I7\)\*I\d+,\s*([0-9.]+)\)")
_LEADING = re.compile(r"^\s*([0-9][0-9,.]*)")


def ap_sprinkler(b: Book, q: str, calc: str, nozzle: str) -> tuple[dict, dict]:
    """FS-063 rule 6. 'Unit Cost & Calculation'!J8 = MIN(MIN(I6,I7)*I8,19000):
    the lower of the unit cost and the cost before GST, times the percent, under a
    ceiling; K8 = I7 - J8 + I28, the farmer paying half the GST (I28 = I23*2.5 %)."""
    formula = str(b.f[calc]["J8"].value)
    m = _CEILING.search(formula.replace(" ", ""))
    assert m, formula
    assert str(b.f[q]["I28"].value).replace(" ", "") == "=+I23*0.025", b.f[q]["I28"].value
    sys_ = system("field_inspection_only", "size_area_2d", "area_matrix", "excl_gst", [],
                  ["field", "transport"])
    params = {"pipe_size_from_input": "1", "share_floor_zero": "0", "inspection_floor": "0",
              "inspection_rate": "0", "gst_material_half": "0.025", "gst_service_half": "0.09",
              "exact_area_match_required": "0", "subsidy_cap": s(Decimal(m.group(1))),
              "farmer_gst_fraction": "0.5", "transport_flat": "1"}
    cats = [cat("small", "Small farmer (50%)", int(b.num(calc, "I8") * 100), 1),
            cat("other", "Other farmer (50%)", int(b.num(calc, "I9") * 100), 2)]
    pub = {"scheme": "AP", "system_type": "sprinkler", "system": sys_, "parameters": params,
           "categories": cats, "unit_cost": sprinkler_unit_cost(b)}
    quantities = sprinkler_quantities(b, nozzle)
    # J22 'Fix': transport is one flat charge per quotation, not per hectare (rule 6)
    assert b.text(q, "F22") == "Fix", b.text(q, "F22")
    quantities["rows"]["transport"] = dict.fromkeys(quantities["areas"], "1")
    chosen = b.text(q, "H10")
    row = next(r for r in (8, 9) if (b.text(calc, f"H{r}") or "").replace(" ", "")
               == (chosen or "").replace(" ", ""))
    inp = example_input(crops=[crop(b.num(q, "H8"), None, [])], nozzle=nozzle,
                        pipe_size_mm=size_mm(b.raw(q, "H7")))
    sample = {"component_rates": sprinkler_rates(b), "quantities": quantities,
              "example": {"input": inp, "category": cats[row - 8]["code"],
                          "expected": expected(b.num(q, "I25"), b.num(calc, "I6"),
                                               b.num(calc, f"J{row}"), b.num(calc, f"K{row}"))}}
    return pub, sample


def _manual(b: Book, sheet: str, cell: str) -> Decimal:
    """'134367 (MANUAL)': the number the portal printed, typed into the sheet."""
    m = _LEADING.match(str(b.raw(sheet, cell)))
    assert m, (cell, b.raw(sheet, cell))
    return Decimal(m.group(1).replace(",", ""))


def ap_drip(b: Book) -> tuple[dict, dict]:
    """FS-063 rules 7 to 11. The portal's BOQ for one farmer: only the lines and the
    figures are copied, never the name, ids, survey or Aadhaar (rule 15)."""
    sh = "MIP_GET_VIEW_BOQDripSystem"

    def block(rows: range, desc: str, size: str, rate: str, qty: str) -> list[dict[str, str]]:
        out = []
        for r in rows:
            d, rv = b.text(sh, f"{desc}{r}"), b.raw(sh, f"{rate}{r}")
            if not d or rv is None or isinstance(rv, str):
                continue
            label = " ".join(f"{d.strip()} {(b.text(sh, f'{size}{r}') or '').strip()}".split())
            out.append({"description": label, "uom": "No.", "rate": s(b.num(sh, f"{rate}{r}")),
                        "qty": s(b.num(sh, f"{qty}{r}", Decimal(0)))})
        return out

    head = block(range(23, 32), "D", "N", "W", "Z")
    field = (block(range(36, 43), "F", "N", "W", "Z") + block(range(47, 55), "F", "N", "W", "Z")
             + block(range(58, 64), "G", "O", "X", "AA"))
    pct = _manual(b, sh, "AE77")
    pub = {"scheme": "AP", "system_type": "drip",
           "system": system("portal_manual", "none", "manual", "excl_gst", [], [], head=True),
           "parameters": {"farmer_gst_fraction": "0.5", "subsidy_round_rupee": "1",
                          "gst_inclusive_rate": "0.05"},
           "categories": [cat(f"eligible_{s(pct)}", f"Eligible {s(pct)}% (portal)", int(pct), 1)]}
    inp = example_input(crops=[crop(dec(b.raw(sh, "M5")), b.num(sh, "M10"), field)],
                        head_lines=head, guided_cost=s(_manual(b, sh, "AE78")),
                        fittings_amount=s(b.num(sh, "AF70")))
    sample = {"component_rates": [], "quantities": None,
              "example": {"input": inp, "category": pub["categories"][0]["code"],
                          "expected": expected(b.num(sh, "L72"), _manual(b, sh, "AE78"),
                                               _manual(b, sh, "K79"), _manual(b, sh, "AE79")),
                          # the portal prints whole rupees (rule 4; the share is compared to the rupee)
                          "rupees": True}}
    return pub, sample


# ── CG drip (GST-inclusive, label x anchors) ─────────────────────────────────

def cg_drip(b: Book) -> tuple[dict, dict]:
    q, calc = "DRIP Quotation", "Unit cost and Calculation"
    anchors = {"0.4": "B", "1": "N", "2": "AH", "3": "BB", "4": "BV", "5": "CP"}
    for a, col in anchors.items():
        assert area_key(b.raw(calc, f"{col}4")) == a, (col, b.raw(calc, f"{col}4"))
    rows = {}
    for r in range(5, 20):
        label = b.raw(calc, f"A{r}")
        key = s(dec(label)) if isinstance(label, (int, float)) else str(label)
        rows[key] = {a: s(b.num(calc, f"{col}{r}")) for a, col in anchors.items()}
    pub = {"scheme": "CG", "system_type": "drip",
           "system": system("gst_inclusive", "label_anchor_linear", "manual", "incl_gst", [], []),
           "parameters": {"fittings_rate": s(b.num(q, "F35")), "flat_inspection": s(b.num(q, "G37")),
                          "unit_cost_round_rupee": "1", "gst_inclusive_rate": "0.05"},
           "categories": [cat("small", "Small farmer (55%)", 55, 1), cat("big", "Big farmer (45%)", 45, 2)],
           "unit_cost": {"kind": "rows", "rows": rows}}
    inp = example_input(crops=[crop(b.num(q, "G11"), b.num(q, "G9"),
                                    lines_of(b, q, range(16, 34), desc="D", uom="F", rate="G", qty="E"))],
                        spacing_label=b.text(q, "G8"))
    assert b.text(q, "D13") == "Small Farmer (55%)"
    sample = {"component_rates": [], "quantities": None,
              "example": {"input": inp, "category": "small",
                          "expected": expected(b.num(calc, "B24"), b.num(calc, "B23"),
                                               b.num(calc, "C26"), b.num(calc, "D26"))}}
    return pub, sample


# ── RJ drip and mini sprinkler (unit cost only) ──────────────────────────────

RJ_CATS = [cat("small_marginal", "Small or marginal farmer (75%)", 75, 1), cat("sc", "SC farmer (75%)", 75, 2),
           cat("st", "ST farmer (75%)", 75, 3), cat("woman", "Woman farmer (75%)", 75, 4),
           cat("others", "Others farmer (70%)", 70, 5)]


def rj(b: Book, system_type: str) -> tuple[dict, dict]:
    drip = system_type == "drip"
    q = "DRIP QTY" if drip else "Mini Sprinkler QTY"
    uc = "Unit Cost"
    table_rows = range(3, 19) if drip else range(5, 7)
    cols = "BCDEFGH"
    rows = {}
    for r in table_rows:
        label = b.text(uc, f"A{r}")
        cells = {area_key(b.raw(uc, f"{c}2")): s(b.num(uc, f"{c}{r}")) for c in cols
                 if b.raw(uc, f"{c}{r}") is not None}
        if label and cells:
            rows[label.strip()] = cells
    pub = {"scheme": "RJ", "system_type": system_type,
           "system": system("unit_cost_only", "label_anchor_linear", "manual", "unit_cost_only", [], [],
                            head=True),
           "parameters": {"min_area_prorate": "1", "zero_above_max_area": "1", "head_unit_deduction": "8000",
                          "gst_inclusive_rate": "0.05"},
           "categories": RJ_CATS, "unit_cost": {"kind": "rows", "rows": rows}}

    # add-on and deduction tables: hydrocyclone A:B, gravel C:D, PVC A:D
    if drip:
        hyd, grav, pvc, cats_at, total_cell, final_cell = range(28, 31), range(28, 32), range(35, 38), 31, "K56", "K29"
    else:
        hyd, grav, pvc, cats_at, total_cell, final_cell = range(12, 15), range(12, 16), range(19, 22), 15, "K61", "K13"
    rates: list[dict[str, Any]] = []
    for r in hyd:
        rates.append({"code": "hydrocyclone", "description": f"Hydrocyclone filter {s(b.num(uc, f'A{r}'))} m3",
                      "uom": "No.", "rate": s(b.num(uc, f"B{r}")), "pipe_size_mm": int(b.num(uc, f"A{r}")), "nozzle": None})
    for r in grav:
        rates.append({"code": "gravel_filter", "description": f"Gravel filter {s(b.num(uc, f'C{r}'))} m3",
                      "uom": "No.", "rate": s(b.num(uc, f"D{r}")), "pipe_size_mm": int(b.num(uc, f"C{r}")), "nozzle": None})
    for r in pvc:
        qty, rate = b.num(uc, f"B{r}"), b.num(uc, f"C{r}")
        rates.append({"code": "pvc_deduction", "description": f"{b.text(uc, f'A{r}')}, {s(qty)} m at {s(rate)}",
                      "uom": "Rs", "rate": s(qty * rate), "pipe_size_mm": int(re.match(r"\d+", b.text(uc, f"A{r}")).group()), "nozzle": None})

    head = lines_of(b, q, range(19, 29), desc="B", uom="G", rate="I", qty="H")
    field = lines_of(b, q, range(31, 60 if drip else 65), desc="B", uom="G", rate="I", qty="H")
    inp = example_input(crops=[crop(b.num(q, "K8"), b.num(q, "K10"), field)], head_lines=head,
                        spacing_label=b.text(q, "K12"), hydrocyclone_m3=s(b.num(q, "K13")),
                        gravel_filter_m3=s(b.num(q, "K14")),
                        head_unit_existing=b.text(q, "C13") == "Yes", pvc_existing=b.text(q, "C14") == "Yes")
    chosen = b.text(q, "C12")
    row = next(cats_at + i for i in range(5) if b.text(uc, f"J{cats_at + i}") == chosen)
    sample = {"component_rates": rates, "quantities": None,
              "example": {"input": inp, "category": RJ_CATS[row - cats_at]["code"],
                          "expected": expected(b.num(q, total_cell), b.num(uc, final_cell),
                                               b.num(uc, f"M{row}"), b.num(uc, f"N{row}"))}}
    return pub, sample


def main() -> int:
    data = next((p for p in DATA_CANDIDATES if p.is_dir()), None)
    if data is None:
        print("quotation formats not found", file=sys.stderr)
        return 1
    states, samples = {}, {}
    for key, rel in FILES.items():
        b = Book(data / rel)
        state, system_type = key.split("/")
        if key == "AP/drip":
            pub, smp = ap_drip(b)
        elif key == "UP/drip":
            pub, smp = up_drip(b)
        elif key == "UP/mini_sprinkler":
            pub, smp = up_mini(b)
        elif system_type == "sprinkler":
            pub, smp = sprinkler(b, state)
        elif key == "CG/drip":
            pub, smp = cg_drip(b)
        else:
            pub, smp = rj(b, system_type)
        pub["source"] = rel
        smp["source"] = rel
        states[key], samples[key] = pub, smp
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "states.json").write_text(json.dumps(states, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (OUT / "states_samples.json").write_text(json.dumps(samples, indent=2, ensure_ascii=False) + "\n",
                                             encoding="utf-8")
    print(f"wrote {len(states)} entries to states.json and states_samples.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
