"""List exports to Excel (FS-030): drain a list service's pages, then write a workbook.

The export never has a query of its own. Each endpoint hands `drain()` a closure
over the list service the screen calls, with the caller's session, so the service
predicate and RLS both apply and a field officer who sees forty leads downloads
forty (Build-Plan 9.3). Reports (FS-02x) reuse `drain()` and `workbook()`.

The writer works from `model_dump(mode="json")`, so every value arrives as JSON:
money as a decimal string, timestamps as ISO strings. The column spec types each
column and the writer turns it back into a cell Excel can sum or sort.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from collections.abc import Awaitable, Callable, Sequence
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Any, Final, Literal

import structlog
from pydantic import BaseModel

from api.errors import ApiError

# structlog, as the rest of api/: a stdlib logger here was never configured and its
# INFO line was dropped (code review F-1)
log = structlog.get_logger()

IST: Final = dt.timezone(dt.timedelta(hours=5, minutes=30))
MAX_ROWS: Final = 5000             # GAP-301: refuse above, never truncate
PAGE: Final = 100                  # every list service clamps to 100
CELL_MAX: Final = 32767            # Excel's own limit; it drops a longer cell on "repair"
XLSX_TYPE: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# what openpyxl refuses to store (its own ILLEGAL_CHARACTERS_RE, restated so the
# rule is visible here): C0 controls except tab, newline and carriage return
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

Kind = Literal["text", "money", "decimal", "int", "datetime", "date", "bool", "names"]


class ExportTooLargeError(ApiError):
    status_code = 422
    code = "export_too_large"


class FilterNotExportableError(ApiError):
    status_code = 422
    code = "filter_not_exportable"


@dataclass(frozen=True)
class Column:
    """One column: its header, where the value sits in the item (dotted), and how to
    write it. `names` joins a list of objects by their `name`."""

    header: str
    path: str
    kind: Kind = "text"


Page = tuple[Sequence[BaseModel], str | None]


async def drain(fetch: Callable[[str | None], Awaitable[Page]], *,
                cap: int | None = None) -> list[dict[str, Any]]:
    """Every row the list would show, across all pages, as JSON dicts.

    `fetch(cursor)` returns one page and the next cursor (None on the last). Stops
    as soon as the count passes `cap`, so a too-large export costs one page past the
    limit, not the whole table.
    """
    cap = MAX_ROWS if cap is None else cap
    rows: list[dict[str, Any]] = []
    cursor: str | None = None
    while True:
        items, cursor = await fetch(cursor)
        rows.extend(i.model_dump(mode="json") for i in items)
        if len(rows) > cap:
            raise ExportTooLargeError(
                f"This list has more than {cap} rows. Narrow the filters and try again.")
        if not cursor:
            return rows


def _get(item: dict[str, Any], path: str) -> Any:
    value: Any = item
    for part in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def clean_text(value: str) -> str:
    """Strip what Excel cannot store, and cut what it would drop."""
    out = _ILLEGAL.sub("", value)
    if len(out) > CELL_MAX:
        out = out[:CELL_MAX - 1] + "…"
    return out


def _to_ist(value: str) -> dt.datetime | None:
    # every exported timestamp is a timestamptz serialised with its offset; a naive
    # string would be read as UTC (code review F-2)
    try:
        parsed = dt.datetime.fromisoformat(value)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=dt.UTC)
    # Excel stores no zone: write the IST wall time (rule 12)
    return parsed.astimezone(IST).replace(tzinfo=None)


def cell_value(value: Any, kind: Kind) -> Any:
    """The Python value for one cell, or None for an empty one."""
    if value is None or value == "":
        return None
    if kind in ("money", "decimal"):
        try:
            return Decimal(str(value))      # openpyxl writes Decimal as a number (rule 4)
        except InvalidOperation:
            return clean_text(str(value))
    if kind == "int":
        return int(value)
    if kind == "bool":
        return "Yes" if value else "No"
    if kind == "datetime":
        return _to_ist(str(value)) or clean_text(str(value))
    if kind == "date":
        try:
            return dt.date.fromisoformat(str(value)[:10])
        except ValueError:
            return clean_text(str(value))
    if kind == "names":
        if not isinstance(value, list):
            return None
        names = [str(v.get("name", "")) if isinstance(v, dict) else str(v) for v in value]
        return clean_text(", ".join(n for n in names if n)) or None
    return clean_text(str(value))


_FORMATS: Final[dict[str, str]] = {"money": "#,##0.00", "datetime": "yyyy-mm-dd hh:mm",
                                   "date": "yyyy-mm-dd"}


def workbook(title: str, columns: Sequence[Column], rows: Sequence[dict[str, Any]]) -> bytes:
    """One sheet: a frozen, filterable header row, then one row per item.

    Every text cell is forced to the string type after assignment. openpyxl reads a
    value starting with `=` as a formula, and `append()` does the same, so without
    this a farmer named `=HYPERLINK(...)` becomes a live link (rule 9).
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = title[:31]
    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=col.header)
        cell.data_type = "s"
    for r, item in enumerate(rows, start=2):
        for c, col in enumerate(columns, start=1):
            value = cell_value(_get(item, col.path), col.kind)
            if value is None:
                continue
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            fmt = _FORMATS.get(col.kind)
            if fmt and not isinstance(value, str):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    if columns:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    for c, col in enumerate(columns, start=1):
        width = 20 if col.kind in ("datetime", "money") else max(10, min(40, len(col.header) + 4))
        ws.column_dimensions[get_column_letter(c)].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filename(stem: str, today: dt.date | None = None) -> str:
    day = today or dt.datetime.now(IST).date()
    return f"{stem}-{day.isoformat()}.xlsx"


def log_export(user_id: str, stem: str, filters: dict[str, Any], count: int) -> None:
    """One line per export: who, which list, which filters were set (names only, as
    `q` can hold a mobile number: rule 13), and how many rows left."""
    used = sorted(k for k, v in filters.items() if v not in (None, "", False))
    log.info("export", list=stem, user=user_id, filters=used, rows=count)


# ── the columns of each list (FS-030 6.1) ────────────────────────────────────

LEADS: Final = (
    Column("Inquiry no", "inquiry_no"), Column("Stage", "stage"), Column("Priority", "priority"),
    Column("Score", "score", "decimal"), Column("Farmer name", "farmer_name"),
    Column("Mobile", "mobile"), Column("Email", "email"), Column("Village", "village"),
    Column("Territory", "territory.name"), Column("Inquiry type", "inquiry_type"),
    Column("Source", "source"), Column("MIS system", "mis_system"),
    Column("Crops", "crops", "names"), Column("Land (acres)", "land_acres", "decimal"),
    Column("Estimated value", "estimated_value", "money"), Column("Owner", "owner.full_name"),
    Column("Office", "owner_org_unit.name"), Column("Partner", "assigned_partner.name"),
    Column("Lost reason", "lost_reason.name"), Column("Lost note", "lost_note"),
    Column("First contacted", "first_contacted_at", "datetime"),
    Column("Last activity", "last_activity_at", "datetime"),
    Column("Created", "created_at", "datetime"), Column("Created by", "created_by.full_name"),
)

_TOTALS: Final = (
    Column("Subtotal", "totals.gross", "money"), Column("Discount", "totals.discount", "money"),
    Column("Taxable", "totals.taxable", "money"), Column("CGST", "totals.cgst", "money"),
    Column("SGST", "totals.sgst", "money"), Column("IGST", "totals.igst", "money"),
    Column("Total", "totals.total", "money"),
)

QUOTATIONS: Final = (
    Column("Quote no", "quote_no"), Column("Version", "version", "int"),
    Column("Status", "status"), Column("Sales type", "sales_type"),
    Column("Party", "party_name"), Column("Party mobile", "party_mobile"),
    Column("Lead", "lead.inquiry_no"), Column("Partner", "partner.name"),
    Column("Owner", "owner.full_name"), *_TOTALS,
    Column("Provisional", "is_provisional", "bool"), Column("Valid until", "valid_until", "date"),
    Column("Sent", "sent_at", "datetime"), Column("Viewed", "viewed_at", "datetime"),
    Column("Awaiting approval", "awaiting_approval", "bool"),
    Column("Superseded by version", "superseded_by.version", "int"),
    Column("Created", "created_at", "datetime"),
)

ORDERS: Final = (
    Column("Order no", "order_no"), Column("Status", "status"), Column("Order type", "order_type"),
    Column("Party", "party_name"), Column("Partner", "partner.name"),
    Column("Owner", "owner.full_name"), *_TOTALS,
    Column("Provisional", "is_provisional", "bool"),
    Column("Dispatched %", "dispatched_pct", "int"),
    Column("Waiting on", "approval_waiting_on"), Column("Submitted", "submitted_at", "datetime"),
    Column("Created", "created_at", "datetime"),
)

COMPLAINTS: Final = (
    Column("Complaint no", "complaint_no"), Column("Status", "status"),
    Column("Type", "complaint_type.name"), Column("Severity", "severity"),
    Column("Contact", "contact_name"), Column("Partner", "partner.name"),
    Column("Owner", "owner.full_name"), Column("Submitted", "first_submitted_at", "datetime"),
    Column("SLA breached", "breached", "bool"), Column("Created", "created_at", "datetime"),
)

TASKS: Final = (
    Column("Title", "title"), Column("Type", "task_type"), Column("Status", "status"),
    Column("Overdue", "overdue", "bool"), Column("Due", "due_at", "datetime"),
    Column("Assigned to", "assigned_to.full_name"), Column("Assigned by", "assigned_by.full_name"),
    Column("Lead", "lead.inquiry_no"), Column("Partner", "partner.name"),
    Column("Order", "sales_order.order_no"), Column("Outcome", "outcome"),
    Column("Completed", "completed_at", "datetime"), Column("Created", "created_at", "datetime"),
)

SUBSIDY_APPLICATIONS: Final = (
    Column("Application no", "application_no"), Column("Reg no", "reg_no"),
    Column("Status", "status"), Column("Stage", "current_stage.name"),
    Column("Scheme", "scheme"), Column("System", "system_type"),
    Column("Category", "category.name"), Column("Farmer", "farmer_name"),
    Column("Mobile", "mobile"), Column("Village", "village"), Column("Survey no", "survey_no"),
    Column("Territory", "territory.name"), Column("Partner", "partner.name"),
    Column("Area (ha)", "total_area", "decimal"),
    Column("Total cost", "figures.total_cost", "money"),
    Column("Subsidy", "figures.subsidy", "money"),
    Column("Farmer share", "figures.farmer_share", "money"),
    Column("Owner", "owner.full_name"), Column("Office", "owner_org_unit.name"),
    Column("Days at stage", "ageing.days_in_stage", "int"),
    Column("Full FP received", "full_fp_received_on", "date"),
    Column("Created", "created_at", "datetime"),
)

PARTNERS: Final = (
    Column("Code", "code"), Column("Name", "name"), Column("Type", "partner_type"),
    Column("Price tier", "price_tier"), Column("Parent", "parent.name"),
    Column("Territory", "territory.name"), Column("Contact", "contact_name"),
    Column("Mobile", "mobile"), Column("Email", "email"), Column("GSTIN", "gstin"),
    Column("GST registered", "is_gst_registered", "bool"),
    Column("Credit limit", "credit_limit", "money"),
    Column("Payment terms (days)", "payment_terms_days", "int"),
    Column("Active", "is_active", "bool"), Column("Users", "users", "int"),
    Column("Created", "created_at", "datetime"),
)

USERS: Final = (
    Column("Name", "full_name"), Column("Type", "user_type"), Column("Role", "role.name"),
    Column("Office", "org_unit.name"), Column("Partner", "partner.name"),
    Column("Email", "email"), Column("Mobile", "mobile"), Column("Active", "is_active", "bool"),
    Column("Last sign-in", "last_login_at", "datetime"), Column("Open leads", "open_leads", "int"),
)

DEALER_COMMISSIONS: Final = (
    Column("Application no", "application.application_no"), Column("Reg no", "application.reg_no"),
    Column("Dealer", "partner.name"), Column("Status", "status"),
    Column("Cost excl. GST", "cost_excl_gst", "money"), Column("GI fitting", "gi_fitting", "money"),
    Column("PVC-HDPE fitting", "pvc_hdpe_fitting", "money"),
    Column("Installation", "installation", "money"),
    Column("Commission base", "commission_base", "money"),
    Column("Commission %", "commission_pct", "decimal"),
    Column("Commission", "commission_amount", "money"), Column("TOD base", "tod_base", "money"),
    Column("TOD %", "tod_pct", "decimal"), Column("TOD", "tod_amount", "money"),
    Column("Total", "total", "money"), Column("Recorded by", "recorded_by.name"),
    Column("Recorded", "recorded_at", "datetime"), Column("Paid on", "paid_on", "date"),
    Column("Payment reference", "payment_reference"),
)

MARKETING_ORDERS: Final = (
    Column("Order no", "order_no"), Column("Status", "status"), Column("Partner", "partner.name"),
    Column("Requested by", "requested_by.name"), Column("Office", "office.name"),
    Column("Value", "totals.value", "money"),
    Column("Company share", "totals.company_share", "money"),
    Column("Dealer share", "totals.dealer_share", "money"),
    Column("Provisional", "is_provisional", "bool"), Column("Created", "created_at", "datetime"),
)

_AGE = ("today_to_supply", "inward_to_submission", "wo_to_tpa_received",
        "tpa_cleared_to_inspection_sent", "inspection_sent_to_tr", "fp_submitted_to_full_fp")
SUBSIDY_AGEING: Final = (
    Column("Application no", "application.application_no"), Column("Reg no", "application.reg_no"),
    Column("Farmer", "application.farmer_name"), Column("District", "application.district"),
    Column("Status", "application.status"), Column("Stage", "application.stage"),
    *(Column(name.replace("_", " ").capitalize() + " (days)", f"{name}.days", "int")
      for name in _AGE),
)
SUBSIDY_STAGES: Final = (
    Column("Stage", "seq", "int"), Column("Name", "name"), Column("Applications", "count", "int"),
    Column("Total cost", "total_cost", "money"), Column("Subsidy", "subsidy", "money"),
    Column("Farmer share", "farmer_share", "money"),
    Column("Oldest days in stage", "oldest_days_in_stage", "int"),
)
SUBSIDY_SUPPLY: Final = (
    Column("District", "district"), Column("Supplied", "supplied", "int"),
    Column("Not supplied", "not_supplied", "int"),
    Column("Supplied cost", "supplied_cost", "money"),
    Column("Not supplied cost", "not_supplied_cost", "money"),
)
