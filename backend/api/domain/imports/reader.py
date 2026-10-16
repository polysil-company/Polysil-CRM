"""Reading an uploaded sheet into header text and rows of cell text (FS-062).

No database. The service stores what this returns and checks it later, so every
cell leaves here as a string: the one form both an `.xlsx` and a `.csv` can give,
and the one the person sees on the matching screen.

The hazards are in the file, not the data (FS-062 EC-8, EC-15, EC-16, EC-20):

- A 10 MB `.xlsx` can unzip to gigabytes. The zip directory's sizes are summed
  before openpyxl opens anything.
- A sheet formatted down to row 1,048,576 yields a million blank rows. Blank rows
  are skipped, and reading stops at a physical-row ceiling.
- Excel's "Unicode Text" export is UTF-16 with a byte-order mark, tab-separated.
"""

from __future__ import annotations

import codecs
import csv
import io
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field

from api.domain.imports.cells import cell_text, clean_text

MAX_DATA_ROWS = 5000
MAX_PHYSICAL_ROWS = 50_000
MAX_COLUMNS = 60
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024
MAX_CELL_CHARS = 1000
HEADER_SCAN = 10         # non-blank rows searched for the header (FS-067 rule 1)
MAX_LABEL_CHARS = 20
MAX_NAME_CHARS = 200

Score = Callable[[Sequence[str]], int]


class ImportUnreadableError(ValueError):
    """The file is not one this reads. `code` is the API's error code."""

    def __init__(self, message: str, code: str = "import_unreadable") -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class SheetRow:
    row: int                 # the row number in the person's sheet, 1-based
    cells: tuple[str, ...]


@dataclass
class Sheet:
    sheets: list[str]
    sheet: str | None        # None for a CSV
    header_row: int
    headers: list[str]
    rows: list[SheetRow] = field(default_factory=list)


def read_upload(data: bytes, filename: str, *, sheet: str | None = None,
                header_row: int | None = None, score: Score | None = None) -> Sheet:
    """An `.xlsx` or a `.csv`, judged by content first and the name second.

    `score` rates a row as a header: how many of an import kind's column names it
    holds. Without it, or with `header_row`, the header is found as before."""
    if data[:4] == b"PK\x03\x04":
        return _read_xlsx(data, sheet=sheet, header_row=header_row, score=score)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ImportUnreadableError("This is an old .xls file, or a password-protected workbook. "
                               "Open it in Excel and save it as .xlsx.")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        raise ImportUnreadableError("This .xlsx file is damaged. Open it in Excel and save it "
                                    "again.")
    if name.endswith(".xls"):
        raise ImportUnreadableError("Old .xls files are not read. Open it in Excel and save it as "
                               ".xlsx.")
    return _read_csv(data, header_row=header_row, score=score)


# ── xlsx ─────────────────────────────────────────────────────────────────────

def _read_xlsx(data: bytes, *, sheet: str | None, header_row: int | None,
               score: Score | None) -> Sheet:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise ImportUnreadableError("This workbook unpacks to more than 100 MB. Save only "
                                       "the sheet you need as a new file.")
            names = set(z.namelist())
    except zipfile.BadZipFile as exc:
        raise ImportUnreadableError("This .xlsx file is damaged. Open it in Excel and save it "
                               "again.") from exc
    if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
        raise ImportUnreadableError("This is a zip file, not an Excel workbook.")

    import openpyxl  # heavy; only an upload pays for it

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a dozen types for a broken file
        raise ImportUnreadableError("This workbook could not be read. Open it in Excel and save it "
                               "again.") from exc
    try:
        visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
        if not visible:
            raise ImportUnreadableError("This workbook has no visible sheet.")
        if sheet is None:
            title = visible[0]
        elif sheet in wb.sheetnames:
            title = sheet
        else:
            raise ImportUnreadableError(f"No sheet named {sheet!r}.", code="validation_error")
        ws = wb[title]
        raw: list[tuple[int, list[str]]] = []
        filled = 0
        for n, values in enumerate(ws.iter_rows(), start=1):
            if n > MAX_PHYSICAL_ROWS:
                break
            cells = [cell_text(getattr(c, "value", None), getattr(c, "number_format", None))
                     for c in values[:MAX_COLUMNS]]
            raw.append((n, cells))
            filled += any(cells)
            if filled > _read_ceiling(header_row):
                break
    finally:
        wb.close()
    return _assemble(raw, sheets=list(wb.sheetnames), sheet=title, header_row=header_row,
                     score=score)


def _read_ceiling(header_row: int | None) -> int:
    """Stop reading once this many non-blank rows are in: one past the data cap,
    plus the header and any title rows above it. `_assemble` applies the exact cap.
    A header still to be found may be any of the first `HEADER_SCAN` rows (FS-067
    rule 1a): a lower one must not shrink the room for the over-cap check."""
    return MAX_DATA_ROWS + 1 + (header_row or HEADER_SCAN)


# ── csv ──────────────────────────────────────────────────────────────────────

def _decode(data: bytes) -> str:
    if data.startswith(codecs.BOM_UTF8):
        return data[len(codecs.BOM_UTF8):].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    if b"\x00" in data[:4096]:
        raise ImportUnreadableError("This file is not text. Upload an .xlsx or a .csv.")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _read_csv(data: bytes, *, header_row: int | None, score: Score | None) -> Sheet:
    text = _decode(data)
    sample = text[:16384]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
    raw: list[tuple[int, list[str]]] = []
    filled = 0
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        for n, values in enumerate(reader, start=1):
            if n > MAX_PHYSICAL_ROWS:
                break
            cells = [clean_text(v)[:MAX_CELL_CHARS] for v in values[:MAX_COLUMNS]]
            raw.append((n, cells))
            filled += any(cells)
            if filled > _read_ceiling(header_row):
                break
    except csv.Error as exc:
        raise ImportUnreadableError(f"This CSV could not be read: {exc}.") from exc
    return _assemble(raw, sheets=[], sheet=None, header_row=header_row, score=score)


# ── both ─────────────────────────────────────────────────────────────────────

def _assemble(raw: list[tuple[int, list[str]]], *, sheets: list[str], sheet: str | None,
              header_row: int | None, score: Score | None = None) -> Sheet:
    labelled = False
    if header_row is None:
        filled = [n for n, cells in raw if any(cells)]
        if not filled:
            raise ImportUnreadableError("This sheet is empty.")
        header_row = _best_header(raw, filled, score)
        labelled = header_row != filled[0]
    by_row = dict(raw)
    if header_row not in by_row or not any(by_row[header_row]):
        raise ImportUnreadableError(f"Row {header_row} is empty, so it cannot be the header row.",
                               code="validation_error")
    header_cells = by_row[header_row]
    width = max([len(header_cells)] + [len(c) for n, c in raw if n > header_row])
    width = min(width, MAX_COLUMNS)
    labels = _labels(raw, header_row, width) if labelled else [""] * width
    headers = _header_names(header_cells, width, labels)
    rows = []
    for n, cells in raw:
        if n <= header_row or not any(cells):
            continue
        padded = (list(cells) + [""] * width)[:width]
        rows.append(SheetRow(row=n, cells=tuple(padded)))
    if len(rows) > MAX_DATA_ROWS:
        raise ImportUnreadableError(f"This sheet has more than {MAX_DATA_ROWS:,} rows. Split it "
                                    f"into smaller files.", code="import_too_many_rows")
    return Sheet(sheets=sheets, sheet=sheet, header_row=header_row, headers=headers, rows=rows)


def _letters(index: int) -> str:
    out = ""
    n = index + 1
    while n:
        n, r = divmod(n - 1, 26)
        out = chr(65 + r) + out
    return out


def _best_header(raw: list[tuple[int, list[str]]], filled: list[int],
                 score: Score | None) -> int:
    """FS-067 rule 1. A row below the first non-blank one is the header only when it
    holds at least two of the kind's column names, and more than the first row
    does: a data cell such as "Nos" or "Distributor" must not win on its own."""
    if score is None:
        return filled[0]
    by_row = dict(raw)
    first = score(by_row[filled[0]])
    best, best_score = filled[0], first
    for n in filled[1:HEADER_SCAN]:
        s = score(by_row[n])
        if s > best_score:
            best, best_score = n, s
    return best if best == filled[0] or (best_score >= 2 and best_score > first) else filled[0]


def _labels(raw: list[tuple[int, list[str]]], header_row: int, width: int) -> list[str]:
    """FS-067 rule 2: short text above the header names its column, as the state codes
    over the client's price columns do. A row's first filled cell is its title and a
    cell with a colon is a note ("Date: 16-10-2026"), so neither is a label."""
    out = [""] * width
    for n, cells in raw:
        if n >= header_row or not any(cells):
            continue
        first = next(i for i, c in enumerate(cells) if c)
        for i, c in enumerate(cells[:width]):
            if c and i != first and len(c) <= MAX_LABEL_CHARS and ":" not in c:
                out[i] = f"{out[i]} {c}".strip()
    return out


def _header_names(cells: list[str], width: int, labels: list[str] | None = None) -> list[str]:
    """A blank header becomes "Column C"; a repeated one gets " (2)" (EC-20). A label
    from above names a blank header, or follows a filled one in brackets."""
    out: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        head = cells[i] if i < len(cells) else ""
        label = labels[i] if labels else ""
        name = f"{head} ({label})" if head and label else head or label or f"Column {_letters(i)}"
        name = name[:MAX_NAME_CHARS]
        key = name.casefold()
        seen[key] = seen.get(key, 0) + 1
        out.append(name if seen[key] == 1 else f"{name} ({seen[key]})")
    return out


__all__ = ["MAX_DATA_ROWS", "ImportUnreadableError", "Sheet", "SheetRow", "read_upload"]
