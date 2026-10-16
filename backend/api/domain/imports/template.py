"""A blank file to fill for each import (FS-068).

Sheet 1 holds only the header row: every field's label, which `suggest` matches
by name. It has no data rows, because an example left in would be imported.
Sheet 2 says what goes in each column, with a made-up example.
"""

from __future__ import annotations

import io

from api.domain.imports.fields import FIELDS, ONE_OF

SHEET_TITLES = {"partners": "Partners", "price_list_items": "Price list",
                "lines": "Order lines"}
FILE_NAMES = {"partners": "partner-import-template.xlsx",
              "price_list_items": "price-list-template.xlsx",
              "lines": "order-lines-template.xlsx"}
HOW_TO = "How to fill"
INSTRUCTION = ("Fill the first sheet. Keep its header row as it is, in row 1. Do not add rows "
               "above it. One row per item. Delete columns you do not use.")
HOW_TO_HEADERS = ("Column", "Needed", "What to put", "Example")


def needed(kind: str, field_key: str) -> str:
    """Yes, No, or the one-of group: "Item code or Description"."""
    specs = {f.key: f for f in FIELDS[kind]}
    if specs[field_key].required:
        return "Yes"
    for group in ONE_OF[kind]:
        if field_key in group:
            return " or ".join(specs[k].label for k in group)
    return "No"


def template(kind: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    specs = FIELDS[kind]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = SHEET_TITLES[kind]
    bold = Font(bold=True)
    for c, spec in enumerate(specs, start=1):
        letter = get_column_letter(c)
        dim = ws.column_dimensions[letter]
        dim.width = max(12, len(spec.label) + 4)
        if spec.text:
            # rule 5: a column style, so a code typed later keeps its leading zeros
            dim.number_format = "@"
        cell = ws.cell(row=1, column=c, value=spec.label)
        cell.font = bold
        if spec.text:
            cell.number_format = "@"
    ws.freeze_panes = "A2"

    how = wb.create_sheet(HOW_TO)
    how.cell(row=1, column=1, value=INSTRUCTION).font = bold
    for c, title in enumerate(HOW_TO_HEADERS, start=1):
        how.cell(row=3, column=c, value=title).font = bold
    for r, spec in enumerate(specs, start=4):
        for c, value in enumerate((spec.label, needed(kind, spec.key), spec.help, spec.example),
                                  start=1):
            cell = how.cell(row=r, column=c, value=value)
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=c == 3, vertical="top")
    for letter, width in zip("ABCD", (22, 24, 60, 28), strict=True):
        how.column_dimensions[letter].width = width
    wb.active = 0

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
